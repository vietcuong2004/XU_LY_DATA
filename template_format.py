"""Copy the template's actual cell and sheet formatting without auto-fit/restyling.

When data blocks grow, extend their corresponding template rows/columns.
Data, formulas, cached results and user corrections are never replaced here.
"""
from copy import copy, deepcopy
from datetime import date, datetime
import re

from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.cell_range import CellRange


def effective_column(sheet, col):
    # Excel stores B:M as one ColumnDimension; looking up C directly is wrong.
    for dim in sheet.column_dimensions.values():
        first = dim.min or 1
        if first <= col <= (dim.max or first):
            return dim
    return None


def week_end(sheet, start):
    row = start
    while re.fullmatch(r'\d{4}/\d{1,2}', str(sheet.cell(row, 1).value or '')):
        row += 1
    if row == start:
        raise ValueError(f'Template {sheet.title}: không có dòng tuần từ dòng {start}.')
    return row - 1


def market_blocks(sheet):
    total = next((c for c in range(2, sheet.max_column+1)
                  if str(sheet.cell(3,c).value or '').strip().upper().startswith('TOTAL ')), None)
    if total is None:
        raise ValueError('Template Breakdown thiếu khối TOTAL KS/KJ.')
    blocks = []
    col = 2
    while col < total:
        area = next((a for a in sheet.merged_cells.ranges if a.min_row == 3 and a.min_col == col), None)
        end = area.max_col if area else col
        blocks.append((col, end))
        col = end+1
    return blocks, total


def map_block(target_start, target_end, ref_start, ref_end):
    result = {}
    for col in range(target_start,target_end+1):
        if col == target_end and target_end > target_start:
            result[col] = ref_end
        else:
            result[col] = min(ref_start+col-target_start, ref_end)
    return result


def apply_template_format(wb, template):
    """Apply the selected template to all three generated sheets, in place."""
    output_blocks, output_total = market_blocks(wb['Breakdown '])
    ref_blocks, ref_total = market_blocks(template['Breakdown '])
    breakdown_map = {1:1}
    for index, (start,end) in enumerate(output_blocks):
        ref_start,ref_end = ref_blocks[min(index,len(ref_blocks)-1)]
        breakdown_map.update(map_block(start,end,ref_start,ref_end))
    breakdown_map.update(map_block(output_total,wb['Breakdown '].max_column,
                                   ref_total,template['Breakdown '].max_column))
    count, ref_count = len(output_blocks),len(ref_blocks)
    for index, (sheet, reference) in enumerate(zip(wb.worksheets,template.worksheets)):
        last_row, last_col = sheet.max_row,sheet.max_column
        start = 14 if index == 0 else 19 if index == 1 else None
        template_end = week_end(reference,start) if start else reference.max_row
        output_end = week_end(sheet,start) if start else last_row

        def ref_row(row):
            if start is None:
                return min(row,reference.max_row)
            if row < start:
                return row
            if row <= output_end:
                if row == output_end and output_end >= template_end:
                    return template_end
                return min(row,template_end-1)
            # Footer rows stay immediately below the resized week block.
            if row > output_end:
                return min(template_end + row - output_end, reference.max_row)

        def ref_col(col,row=None):
            if index == 0:
                return breakdown_map[col]
            if index == 2 or (row is not None and row < 9):
                return col
            if col == 1:
                return 1
            if col <= count+1:
                return min(col,ref_count+1)
            return col-count+ref_count

        # Rebuild header merges using the template's block structure.
        for area in list(sheet.merged_cells.ranges):
            sheet.unmerge_cells(str(area))
        if index == 0:
            areas = [(2,2,2,last_col),(3,3,output_total,last_col),(4,4,output_total,last_col)]
            areas += [(row,row,a,b) for a,b in output_blocks for row in (3,4)]
            for r1,r2,c1,c2 in areas:
                if c2>c1:
                    sheet.merge_cells(start_row=r1,end_row=r2,start_column=c1,end_column=c2)
        else:
            for area in reference.merged_cells.ranges:
                cols = [c for c in range(1,last_col+1) if area.min_col<=ref_col(c,area.min_row)<=area.max_col]
                if cols and area.max_row<=last_row:
                    sheet.merge_cells(start_row=area.min_row,end_row=area.max_row,
                                      start_column=min(cols),end_column=max(cols))

        # Copy styles through the target workbook's style registries (IDs are workbook-local).
        style_cache = {}
        def apply_style(target, source):
            key = tuple(source._style or [])
            if key not in style_cache:
                for attr in ('font','fill','border','alignment','protection'):
                    setattr(target,attr,copy(getattr(source,attr)))
                target.number_format = source.number_format
                target.quotePrefix = source.quotePrefix
                style_cache[key] = copy(target._style)
            else:
                target._style = copy(style_cache[key])

        for row in range(1,last_row+1):
            for col in range(1,last_col+1):
                apply_style(sheet.cell(row,col),reference.cell(ref_row(row),ref_col(col,row)))
        if index == 2:
            # The reference has sparse cutoff dates. New dates in its blank rows
            # keep that row's font/fill/border, using its existing date format
            # rather than displaying an Excel serial number with General.
            date_format = next((c.number_format for cells in reference for c in cells
                                if isinstance(c.value,(date,datetime))), 'yyyy-mm-dd')
            for row in sheet:
                for cell in row:
                    if isinstance(cell.value,(date,datetime)) and cell.number_format == 'General':
                        cell.number_format = date_format
        sheet.row_dimensions.clear()
        for row in range(1,last_row+1):
            source = reference.row_dimensions.get(ref_row(row))
            if source is not None:
                target = copy(source)
                target.index = row
                target.parent = sheet
                apply_style(target,source)
                sheet.row_dimensions[row] = target
        sheet.column_dimensions.clear()
        if index == 2:
            for name,source in reference.column_dimensions.items():
                target = copy(source)
                target.parent = sheet
                apply_style(target,source)
                sheet.column_dimensions[name] = target
        else:
            for col in range(1,last_col+1):
                source = effective_column(reference,ref_col(col))
                if source is not None:
                    target = copy(source)
                    target.index = get_column_letter(col)
                    target.min = target.max = col
                    target.parent = sheet
                    apply_style(target,source)
                    sheet.column_dimensions[target.index] = target
            tail = effective_column(reference,reference.max_column+1)
            if tail is not None:
                target = copy(tail)
                target.index = get_column_letter(last_col+1)
                target.min,target.max = last_col+1,16384
                target.parent = sheet
                sheet.column_dimensions[target.index] = target

        for attr in ('sheet_format','sheet_properties','views','sheet_protection',
                     'page_setup','page_margins','print_options','print_title_rows','print_title_cols'):
            # protection is the Worksheet attribute used by openpyxl.
            if attr == 'sheet_protection':
                sheet.protection = deepcopy(reference.protection)
            else:
                setattr(sheet,attr,deepcopy(getattr(reference,attr)))
        sheet.oddHeader = deepcopy(reference.oddHeader)
        sheet.oddFooter = deepcopy(reference.oddFooter)
        sheet.evenHeader = deepcopy(reference.evenHeader)
        sheet.evenFooter = deepcopy(reference.evenFooter)
        sheet.firstHeader = deepcopy(reference.firstHeader)
        sheet.firstFooter = deepcopy(reference.firstFooter)
        sheet.print_area = ''
        if reference.print_area:
            sheet.print_area = str(reference.print_area).replace(reference.title,sheet.title)
        # Keep template conditional formatting, shifting footer rules below expanded weeks.
        sheet.conditional_formatting = ConditionalFormattingList()
        for group,rules in reference.conditional_formatting._cf_rules.items():
            for area in group.sqref.ranges:
                cols = [c for c in range(1,last_col+1) if area.min_col<=ref_col(c,area.min_row)<=area.max_col]
                if not cols:
                    continue
                delta = output_end-template_end if start and area.min_row>template_end else 0
                target_range = CellRange(min_col=min(cols),max_col=max(cols),
                                         min_row=area.min_row+delta,max_row=area.max_row+delta)
                for rule in rules:
                    cloned = deepcopy(rule)
                    if rule.formula:
                        origin = f'{get_column_letter(area.min_col)}{area.min_row}'
                        dest = f'{get_column_letter(target_range.min_col)}{target_range.min_row}'
                        cloned.formula = [Translator('='+f.lstrip('='),origin=origin).translate_formula(dest)[1:]
                                          for f in rule.formula]
                    sheet.conditional_formatting.add(str(target_range),cloned)
    wb.loaded_theme = template.loaded_theme
