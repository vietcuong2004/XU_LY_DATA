"""Publish a fresh, complete generation after an Excel source edit succeeds."""
from datetime import date, datetime
import math
from pathlib import Path
import re
import shutil
from types import SimpleNamespace
import uuid

import openpyxl
from openpyxl.cell.cell import MergedCell
import excel_engine
import cloud_store


def cell_edit(sheet, address, raw, kind):
    if kind == 'number':
        try:
            raw = float(raw)
        except (ValueError, TypeError) as exc:
            raise ValueError('Nhập số hợp lệ.') from exc
        if not math.isfinite(raw) or abs(raw) > 1e15:
            raise ValueError('Số phải hữu hạn, không vượt 1.000.000.000.000.000.')
    elif kind == 'formula':
        if not isinstance(raw, str) or not raw.startswith('=') or len(raw) > 8192:
            raise ValueError('Công thức phải bắt đầu bằng =, tối đa 8192 ký tự.')
    elif kind == 'date':
        try:
            raw = date.fromisoformat(str(raw)).isoformat()
        except ValueError as exc:
            raise ValueError('Ngày phải có dạng YYYY-MM-DD.') from exc
    elif kind == 'boolean':
        if str(raw).lower() not in ('true', 'false'):
            raise ValueError('Nhập TRUE hoặc FALSE.')
        raw = str(raw).lower() == 'true'
    elif kind == 'blank':
        raw = None
    elif kind == 'text':
        raw = str(raw or '')
        if len(raw) > 32767:
            raise ValueError('Nội dung ô vượt 32767 ký tự.')
    else:
        raise ValueError('Kiểu dữ liệu không hợp lệ.')
    return dict(sheet=sheet, cell=address, kind=kind, value=raw)


def original_edit(sheet, cell):
    value = cell.value
    if cell.data_type == 'f':
        kind = 'formula'
    elif isinstance(value, bool):
        kind = 'boolean'
    elif isinstance(value, (datetime, date)):
        kind, value = 'date', value.isoformat()[:10]
    elif isinstance(value, (int, float)):
        kind = 'number'
    elif value is None:
        kind = 'blank'
    else:
        kind = 'text'
    return cell_edit(sheet, cell.coordinate, value, kind)


def apply(ui, job_id, report, entry, payload):
    if payload.get('revision') != entry['revision']:
        raise ValueError('Phiên vừa thay đổi. Hãy tải lại trước khi sửa tiếp.')
    if not excel_engine.capability()['available']:
        raise ValueError(excel_engine.capability()['message'])
    parent = ui.read_job(job_id)
    if parent.get('superseded_by'):
        raise ValueError('Nguồn đã có phiên mới. Mở phiên mới nhất trong lịch sử để sửa tiếp.')
    folder = ui.directory(job_id)
    original = folder/'baseline'/'input.xlsx'
    if not original.exists():
        original = folder/'input.xlsx'
    current = folder/'input.xlsx' if entry.get('edit_mode') == 'full' else original
    book = openpyxl.load_workbook(current)
    baseline = openpyxl.load_workbook(original)
    pending = payload.get('edits', [payload])
    if not isinstance(pending, list) or not 1 <= len(pending) <= 1000:
        book.close(); baseline.close()
        raise ValueError('Mỗi lần lưu cần từ 1 đến 1000 ô.')
    prepared = []
    try:
        if book.sheetnames == ['SUM'] and any(
                cell.data_type == 'f' and isinstance(cell.value, str) and '!' in cell.value
                for row in book['SUM'] for cell in row):
            raise ValueError('Phiên cũ chỉ chứa SUM, thiếu các sheet liên kết. Hãy tải lại workbook gốc đầy đủ để sửa mọi sheet.')
        seen = set()
        for item in pending:
            index, address = item.get('sheet'), item.get('cell', '')
            if type(index) is not int or not 0 <= index < len(book.worksheets) or not re.fullmatch(r'[A-Z]{1,3}[1-9][0-9]{0,6}', address):
                raise ValueError('Địa chỉ ô không hợp lệ.')
            if (index, address) in seen:
                raise ValueError('Một ô xuất hiện nhiều lần trong yêu cầu lưu.')
            seen.add((index, address))
            sheet = book.worksheets[index]
            max_row, max_column = sheet.max_row, sheet.max_column
            cell = sheet[address]
            if isinstance(cell, MergedCell) or cell.row > max_row or cell.column > max_column or cell.column > 16384 or cell.row > 1048576:
                raise ValueError('Chọn ô đầu của vùng gộp hoặc ô trong bảng dữ liệu.')
            if cell.data_type == 'f' and not isinstance(cell.value, str):
                raise ValueError('Không thể sửa một phần công thức mảng.')
            previous = ui.serial(cell.value)
            sheet_name = sheet.title
            initial_edit = original_edit(sheet_name, baseline[sheet_name][address])
            edit = initial_edit if item.get('restore') else cell_edit(sheet_name, address, item.get('value'), item.get('value_type', 'text'))
            prepared.append((edit, initial_edit, previous, item))
    finally:
        book.close()
        baseline.close()
    # Migrate explicit legacy SUM overrides before applying the requested edit.
    edits = [cell_edit('SUM', key, val, 'blank' if val is None else 'number')
             for key, val in entry.get('overrides', {}).items()]
    edits.extend(p[0] for p in prepared)
    new_id = uuid.uuid4().hex
    new_folder = ui.STORE/new_id
    new_folder.mkdir(parents=True)
    state = {**parent, 'id': new_id, 'created': ui.timestamp(), 'status': 'processing',
             'parent_job': job_id, 'completed': 0, 'total': 0,
             'message': 'Đang tính lại toàn bộ workbook…'}
    state.pop('superseded_by', None)
    ui.write_json(new_folder/'job.json', state)
    if cloud_store.enabled():
        cloud_store.progress(state)
    new_lock = ui.job_lock(new_id)
    new_lock.acquire()
    try:
        shutil.copy2(current, new_folder/'input.xlsx')
        shutil.copy2(folder/'template.xlsx', new_folder/'template.xlsx')
        excel_engine.edit_and_calculate(new_folder/'input.xlsx', edits)
        args = SimpleNamespace(input=str(new_folder/'input.xlsx'), template=str(new_folder/'template.xlsx'),
                               output=str(new_folder/'files'), family=None, season=parent['season'], week=parent['week'],
                               start_week=ui.process.week_arg(parent['start']) if parent.get('start') else None,
                               end_week=ui.process.week_arg(parent['end']) if parent.get('end') else None,
                               date=date.fromisoformat(report['date']) if report.get('date') else None,
                               strict=False, overwrite=False, source_name=parent['source'], defer_publish=True)
        ui.generate(new_id, args)
        if ui.read_job(new_id)['status'] != 'ready':
            raise ValueError(ui.read_job(new_id)['message'])
        updated = ui.get_report(new_id)
        shutil.copy2(original, new_folder/'baseline'/'input.xlsx')
        source = updated['source_file']
        source['edit_mode'] = 'full'
        source['revision'] = entry['revision'] + 1
        source['edits'] = list(entry.get('edits', [])) + [dict(
            time=ui.timestamp(), sheet=edit['sheet'], cell=edit['cell'], before=previous, after=edit['value'],
            kind=edit['kind'], reason=str(item.get('reason', ''))[:500], restored=bool(item.get('restore')))
            for edit, initial_edit, previous, item in prepared]
        changes = dict(entry.get('direct_changes', {}))
        for key, val in entry.get('overrides', {}).items():
            changes['SUM!' + key] = val
        for edit, initial_edit, previous, item in prepared:
            key = edit['sheet'] + '!' + edit['cell']
            if edit == initial_edit:
                changes.pop(key, None)
            else:
                changes[key] = edit
        source['direct_changes'] = changes
        source['changed_cells'] = len(changes)
        source['updated_families'] = [f['id'] for f in updated['families']]
        source['status_label'] = 'edited' if changes else 'matched'
        updated['parent_job'] = job_id
        ui.write_json(new_folder/'review.json', updated)
        if cloud_store.enabled():
            cloud_store.publish(new_folder, ui.read_job(new_id))
        ui.update_job(job_id, superseded_by=new_id)
        if cloud_store.enabled():
            cloud_store.publish(folder, ui.read_job(job_id))
        return {**source, 'new_job_id': new_id}
    except Exception as exc:
        ui.update_job(new_id, status='error', message=str(exc))
        raise
    finally:
        new_lock.release()
