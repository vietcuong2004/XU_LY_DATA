"""Meaningful regression tests for edits, derived values, audit and restore."""
from datetime import date
import json
from pathlib import Path
import shutil
import tempfile
import unittest
from unittest.mock import patch

import openpyxl
from openpyxl.styles import PatternFill,Font

import process
import ui_server as ui


class ReviewTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.patch = patch.object(ui, 'STORE', Path(self.temp.name))
        self.patch.start()
        self.addCleanup(self.patch.stop)
        self.job = 'a' * 32
        self.folder = ui.STORE / self.job
        (self.folder/'files').mkdir(parents=True)
        (self.folder/'baseline').mkdir()
        ui.write_json(self.folder/'job.json', {'id':self.job,'status':'ready'})
        self.entry = {'id':0,'file':'test.xlsx','name':'TEST','sheet':'TEST','items':2,'markets':1,
                      'revision':0,'edits':[],'source_columns':[23,24]}
        wb = openpyxl.Workbook()
        b = wb.active
        b.title = 'Breakdown '
        s = wb.create_sheet('TEST')
        release = wb.create_sheet('Release Qty')
        wb._formula_cache = {}
        b.merge_cells('B3:C3');b['B3']='Italy'
        b.merge_cells('B4:C4');b['B4']='LPIA'
        for c in (2,3):
            b.cell(5,c,'Item '+str(c));b.cell(6,c,'MPG'+str(c));b.cell(7,c,'INTERNATIONAL')
            b.cell(9,c,100);b.cell(10,c,20);b.cell(13,c,'Senft')
            process.sum_formula(b,11,c,[b.cell(9,c),b.cell(10,c)])
        b['A14']='2026/53';b['A15']='2027/53'
        b['B14']=2;b['C14']=3;b['B15']=7;b['C15']='#REF!'
        for r in (9,10,11,14,15):
            process.sum_formula(b,r,4,[b.cell(r,2),b.cell(r,3)])
        for r,src in ((14,9),(15,10),(16,11)):
            process.sum_formula(s,r,2,[b.cell(src,2),b.cell(src,3)])
        for r in (19,20):
            process.sum_formula(s,r,2,[b.cell(r-5,2),b.cell(r-5,3)])
            process.sum_formula(s,r,3,[s.cell(r,2)])
            process.sum_formula(s,r,4,[s.cell(r,3)]+([s.cell(r-1,4)] if r>19 else []))
        s['E3']='NIIGATA';s['E4']='TEST';s['E6']=date(2026,10,5)
        for c,label in enumerate(process.LOGISTICS,3):
            s.cell(9,c,label)
        ui.sync_summary(wb,self.entry)
        process.save_with_cache(wb,self.folder/'files'/'test.xlsx')
        shutil.copy2(self.folder/'files'/'test.xlsx',self.folder/'baseline'/'test.xlsx')
        self.entry.update(ui.analyze(self.folder/'files'/'test.xlsx',self.folder/'baseline'/'test.xlsx',self.entry))
        ui.write_json(self.folder/'review.json',{'families':[self.entry]})

    def edit(self,cell,value,sheet=0,revision=0,**kwargs):
        return ui.apply_edit(self.job,0,dict(cell=cell,value=value,sheet=sheet,revision=revision,**kwargs))

    def values(self):
        wb=openpyxl.load_workbook(self.folder/'files'/'test.xlsx',data_only=True)
        self.addCleanup(wb.close)
        return wb

    def test_quantity_updates_all_sums_and_audit_then_restore(self):
        result=self.edit('B14','12.5',reason='Confirmed')
        wb=self.values()
        self.assertEqual(wb['Breakdown ']['D14'].value,15.5)
        self.assertEqual(wb['TEST']['B19'].value,15.5)
        self.assertEqual(wb['TEST']['D19'].value,15.5)
        self.assertEqual(result['edits'][0]['reason'],'Confirmed')
        self.assertGreater(result['changed_cells'],0)
        result=self.edit('B14',None,revision=1,restore=True)
        self.assertEqual(result['changed_cells'],0)
        self.assertEqual(self.values()['Breakdown ']['B14'].value,2)
        original=openpyxl.load_workbook(self.folder/'baseline'/'test.xlsx',data_only=True)
        self.assertEqual(original['Breakdown ']['B14'].value,2)
        original.close()

    def test_edit_does_not_add_excel_notes_and_preview_keeps_provenance(self):
        report = ui.get_report(self.job)
        report['families'][0]['review_notes'] = {'Breakdown ': {'B14': 'Nguồn: input.xlsx\nSUM!W20'}}
        ui.write_json(self.folder/'review.json', report)
        self.edit('B14', '9')
        wb = openpyxl.load_workbook(self.folder/'files'/'test.xlsx')
        self.assertFalse(any(c.comment for s in wb for row in s for c in row))
        wb.close()
        preview = ui.preview(self.job, 0, 0)
        self.assertIn('SUM!W20', preview['rows'][13][1]['comment'])

    def test_fix_error_and_invalid_week_updates_release_and_cumulative(self):
        result=self.edit('C15','8')
        self.assertFalse(any(p['kind']=='error' for p in result['problems']))
        self.assertTrue(any(p['kind']=='week' for p in result['problems']))
        result=self.edit('A15','2028/01',revision=1)
        self.assertEqual(result['problems'],[])
        wb=self.values()
        self.assertEqual(wb['TEST']['D20'].value,20)
        self.assertEqual(wb['TEST']['A20'].value,'2028/01')
        self.assertEqual(wb['TEST']['E20'].value.date(),date(2028,1,3))
        self.assertEqual(wb['Release Qty']['A4'].value.date(),date(2028,1,3))
        result=self.edit('A15',None,revision=2,restore=True)
        self.assertTrue(any(p['kind']=='week' for p in result['problems']))

    def test_reject_formula_edit_nan_invalid_week_duplicate_and_stale_revision(self):
        for cell,value in [('D14','10'),('B14','NaN'),('B14','=1+1'),('A15','2027/53'),('A15','2026/53'),('B5','new identity')]:
            with self.subTest(cell=cell,value=value),self.assertRaises(ValueError):
                self.edit(cell,value)
        self.edit('B14','5')
        with self.assertRaisesRegex(ValueError,'cập nhật'):
            self.edit('B14','6',revision=0)

    def test_confirm_and_metadata_sync(self):
        self.edit('B9','250')
        wb=self.values()
        self.assertEqual(wb['TEST']['B14'].value,350)
        self.assertEqual(wb['TEST']['B16'].value,390)
        self.edit('B7','NEW',revision=1)
        self.assertEqual(self.values()['TEST']['B11'].value,'NEW / INTERNATIONAL')

    def test_preview_original_and_formula_locks(self):
        self.edit('B14','9')
        result=ui.preview(self.job,0,0)
        cell=result['rows'][13][1]
        self.assertEqual(cell['original'],2)
        self.assertEqual(cell['value'],9)
        self.assertTrue(cell['changed'])
        self.assertEqual(cell['editable'],'number')
        self.assertIsNone(result['rows'][13][3]['editable'])
        self.assertTrue(result['rows'][13][3]['formula'].startswith('=SUM'))

    def test_edit_preserves_template_formatting(self):
        path=self.folder/'files'/'test.xlsx'
        wb=openpyxl.load_workbook(path)
        cell=wb['Breakdown ']['B9']
        cell.fill=PatternFill('solid',fgColor='114433')
        cell.font=Font(name='Arial',size=14,bold=True,color='AA2200')
        cell.number_format='0.0000'
        ui.recalculate(wb)
        process.save_with_cache(wb,path)
        wb.close()
        self.edit('B9','250')
        wb=openpyxl.load_workbook(path)
        cell=wb['Breakdown ']['B9']
        self.assertEqual(cell.fill.fgColor.rgb,'00114433')
        self.assertEqual(cell.font.color.rgb,'00AA2200')
        self.assertEqual(cell.font.size,14)
        self.assertTrue(cell.font.bold)
        self.assertEqual(cell.number_format,'0.0000')
        wb.close()

    def test_text_is_literal_and_date_validation(self):
        self.edit('G19','=not a formula',sheet=1)
        wb=openpyxl.load_workbook(self.folder/'files'/'test.xlsx')
        self.assertEqual(wb['TEST']['G19'].data_type,'s')
        wb.close()
        with self.assertRaises(ValueError):
            self.edit('H19','not a date',sheet=1,revision=1)
        self.edit('H19','2026-10-05',sheet=1,revision=1)
        self.assertEqual(self.values()['TEST']['H19'].value.date(),date(2026,10,5))


if __name__=='__main__':
    unittest.main()
