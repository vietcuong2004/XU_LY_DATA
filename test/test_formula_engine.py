from pathlib import Path
from datetime import datetime
import math
import shutil
import tempfile
import unittest
from zipfile import ZipFile

import openpyxl
from formula_engine import Calculator, UnsupportedFormula, edit_and_calculate

SAMPLE = Path(__file__).resolve().parents[1]/'templates/input_sample.xlsx'


class FormulaTests(unittest.TestCase):
    def test_all_sample_formula_results_match_excel_cached_values(self):
        book = openpyxl.load_workbook(SAMPLE)
        original = openpyxl.load_workbook(SAMPLE, data_only=True)
        self.addCleanup(book.close); self.addCleanup(original.close)
        results = Calculator(book).calculate()
        self.assertGreater(len(results), 67000)
        for (sheet, address), actual in results.items():
            expected = original[sheet][address].value
            if isinstance(actual, (int, float)) and isinstance(expected, (int, float)):
                self.assertTrue(math.isclose(actual, expected, rel_tol=1e-10, abs_tol=1e-8), (sheet, address, actual, expected))
            else:
                self.assertEqual(actual, expected, (sheet, address))

    def test_cross_sheet_edit_keeps_styles_and_all_other_zip_parts(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp)/'input.xlsx'; shutil.copy2(SAMPLE, path)
            edit_and_calculate(path, [dict(sheet='Bel', cell='C16', kind='number', value=321.25)])
            book = openpyxl.load_workbook(path, data_only=True)
            self.assertEqual(book['SUM']['W16'].value, 321.25)
            self.assertEqual(book['SUM']['W2'].value, 'INTERNATIONAL')
            self.assertEqual(book['SUM']['DM132'].value, '#REF!')
            book.close()
            with ZipFile(SAMPLE) as before, ZipFile(path) as after:
                self.assertEqual(before.namelist(), after.namelist())
                for name in before.namelist():
                    if not name.startswith('xl/worksheets/sheet'):
                        self.assertEqual(before.read(name), after.read(name), name)
            before = openpyxl.load_workbook(SAMPLE)
            after = openpyxl.load_workbook(path)
            for a, b in zip(before, after):
                self.assertEqual({k: dict(v) for k,v in a.column_dimensions.items()}, {k: dict(v) for k,v in b.column_dimensions.items()})
                self.assertEqual({k: dict(v) for k,v in a.row_dimensions.items()}, {k: dict(v) for k,v in b.row_dimensions.items()})
                self.assertEqual(a.merged_cells, b.merged_cells)
            before.close(); after.close()

    def test_coercion_error_propagation_and_week_boundaries(self):
        book = openpyxl.Workbook(); s = book.active
        s['A1']='text'; s['A2']=True; s['A3']=2
        s['B1']='=+A1'; s['B2']='=SUM(A1:A3)'; s['B3']='=A1+1'
        s['C1']=datetime(2021, 1, 1); s['C2']='=WEEKNUM(C1,21)'; s['C3']='=WEEKNUM(C1,2)'
        s['D1']='=SUM((A3,A3))'; s['D2']='=1/0'; s['D3']='=D2+1'
        s['E1']='=-2^2'; s['E2']='=2^3^2'; s['E3']='=1<TRUE'
        c=Calculator(book).calculate()
        for address, value in {'B1':'text','B2':2,'B3':'#VALUE!','C2':53,'C3':1,
                               'D1':4,'D2':'#DIV/0!','D3':'#DIV/0!','E1':4,'E2':64,'E3':True}.items():
            self.assertEqual(c[(s.title,address)],value,address)
        book.close()

    def test_unsupported_function_and_cycles_abort_without_modifying_file(self):
        with tempfile.TemporaryDirectory() as temp:
            path=Path(temp)/'input.xlsx'
            book=openpyxl.Workbook();book.active['A1']=3;book.save(path);book.close()
            original=path.read_bytes()
            for formula in ('=INDIRECT("A1")','=A1+1'):
                with self.assertRaises(UnsupportedFormula):
                    edit_and_calculate(path,[dict(sheet='Sheet',cell='A1',kind='formula',value=formula)])
                self.assertEqual(path.read_bytes(),original)

    def test_new_blank_cell_text_date_and_formula_cache_types(self):
        with tempfile.TemporaryDirectory() as temp:
            path=Path(temp)/'input.xlsx'
            book=openpyxl.Workbook();s=book.active
            s['A1']=1;s['D1']='=B1';s['A3']='=B2';s['D3']=1
            book.save(path);book.close()
            edit_and_calculate(path,[dict(sheet='Sheet',cell='B1',kind='text',value='=literal'),
                                     dict(sheet='Sheet',cell='B2',kind='formula',value='="#REF!"')])
            result=openpyxl.load_workbook(path,data_only=True)
            self.assertEqual(result.active['D1'].value,'=literal')
            self.assertEqual(result.active['A3'].value,'#REF!')
            self.assertEqual(result.active['A3'].data_type,'s')
            result.close()

    def test_edit_in_empty_sheet_is_written_and_updates_other_sheet(self):
        with tempfile.TemporaryDirectory() as temp:
            path=Path(temp)/'input.xlsx'
            book=openpyxl.Workbook();book.active.title='Empty'
            book.create_sheet('SUM')['A1']='=Empty!A1*2';book.save(path);book.close()
            edit_and_calculate(path,[dict(sheet='Empty',cell='A1',kind='number',value=7)])
            book=openpyxl.load_workbook(path,data_only=True)
            self.assertEqual(book['Empty']['A1'].value,7)
            self.assertEqual(book['SUM']['A1'].value,14);book.close()

    def test_edit_shared_formula_master_does_not_change_other_formulas(self):
        with tempfile.TemporaryDirectory() as temp:
            original=Path(temp)/'original.xlsx';path=Path(temp)/'input.xlsx'
            book=openpyxl.Workbook();s=book.active
            s['A1']='=B1*2';s['A2']='=B2*2';s['B1']=2;s['B2']=3
            book.save(original);book.close()
            with ZipFile(original) as src,ZipFile(path,'w') as dst:
                for entry in src.infolist():
                    data=src.read(entry.filename)
                    if entry.filename=='xl/worksheets/sheet1.xml':
                        data=data.replace(b'<f>B1*2</f>',b'<f t="shared" si="0" ref="A1:A2">B1*2</f>')
                        data=data.replace(b'<f>B2*2</f>',b'<f t="shared" si="0"/>')
                    dst.writestr(entry,data)
            edit_and_calculate(path,[dict(sheet='Sheet',cell='A1',kind='formula',value='=B1*3')])
            book=openpyxl.load_workbook(path)
            self.assertEqual(book.active['A2'].value,'=B2*2');book.close()
            book=openpyxl.load_workbook(path,data_only=True)
            self.assertEqual(book.active['A1'].value,6)
            self.assertEqual(book.active['A2'].value,6);book.close()


if __name__=='__main__':
    unittest.main()
