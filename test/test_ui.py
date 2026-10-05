"""Meaningful regression tests for edits, derived values, audit and restore."""
from datetime import date
import json
from pathlib import Path
import shutil
import tempfile
import unittest
from unittest.mock import patch

import openpyxl
from openpyxl.styles import PatternFill,Font

import process
import ui_server as ui


class ReviewTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.patch = patch.object(ui, 'STORE', Path(self.temp.name))
        self.patch.start()
        self.addCleanup(self.patch.stop)
        self.job = 'a' * 32
        self.folder = ui.STORE / self.job
        (self.folder/'files').mkdir(parents=True)
        (self.folder/'baseline').mkdir()
        ui.write_json(self.folder/'job.json', {'id':self.job,'status':'ready'})
        self.entry = {'id':0,'file':'test.xlsx','name':'TEST','sheet':'TEST','items':2,'markets':1,
                      'revision':0,'edits':[],'source_columns':[23,24]}
        wb = openpyxl.Workbook()
        b = wb.active
        b.title = 'Breakdown '
        s = wb.create_sheet('TEST')
        release = wb.create_sheet('Release Qty')
        wb._formula_cache = {}
        b.merge_cells('B3:C3');b['B3']='Italy'
        b.merge_cells('B4:C4');b['B4']='LPIA'
        for c in (2,3):
            b.cell(5,c,'Item '+str(c));b.cell(6,c,'MPG'+str(c));b.cell(7,c,'INTERNATIONAL')
            b.cell(9,c,100);b.cell(10,c,20);b.cell(13,c,'Senft')
            process.sum_formula(b,11,c,[b.cell(9,c),b.cell(10,c)])
        b['A14']='2026/53';b['A15']='2027/53'
        b['B14']=2;b['C14']=3;b['B15']=7;b['C15']='#REF!'
        for r in (9,10,11,14,15):
            process.sum_formula(b,r,4,[b.cell(r,2),b.cell(r,3)])
        for r,src in ((14,9),(15,10),(16,11)):
            process.sum_formula(s,r,2,[b.cell(src,2),b.cell(src,3)])
        for r in (19,20):
            process.sum_formula(s,r,2,[b.cell(r-5,2),b.cell(r-5,3)])
            process.sum_formula(s,r,3,[s.cell(r,2)])
            process.sum_formula(s,r,4,[s.cell(r,3)]+([s.cell(r-1,4)] if r>19 else []))
        s['E3']='NIIGATA';s['E4']='TEST';s['E6']=date(2026,10,5)
        for c,label in enumerate(process.LOGISTICS,3):
            s.cell(9,c,label)
        ui.sync_summary(wb,self.entry)
        process.save_with_cache(wb,self.folder/'files'/'test.xlsx')
        shutil.copy2(self.folder/'files'/'test.xlsx',self.folder/'baseline'/'test.xlsx')
        self.entry.update(ui.analyze(self.folder/'files'/'test.xlsx',self.folder/'baseline'/'test.xlsx',self.entry))
        ui.write_json(self.folder/'review.json',{'families':[self.entry]})

    def edit(self,cell,value,sheet=0,revision=0,**kwargs):
        return ui.apply_edit(self.job,0,dict(cell=cell,value=value,sheet=sheet,revision=revision,**kwargs))

    def values(self):
        wb=openpyxl.load_workbook(self.folder/'files'/'test.xlsx',data_only=True)
        self.addCleanup(wb.close)
        return wb

    def test_quantity_updates_all_sums_and_audit_then_restore(self):
        result=self.edit('B14','12.5',reason='Confirmed')
        wb=self.values()
        self.assertEqual(wb['Breakdown ']['D14'].value,15.5)
        self.assertEqual(wb['TEST']['B19'].value,15.5)
        self.assertEqual(wb['TEST']['D19'].value,15.5)
        self.assertEqual(result['edits'][0]['reason'],'Confirmed')
        self.assertGreater(result['changed_cells'],0)
        result=self.edit('B14',None,revision=1,restore=True)
        self.assertEqual(result['changed_cells'],0)
        self.assertEqual(self.values()['Breakdown ']['B14'].value,2)
        original=openpyxl.load_workbook(self.folder/'baseline'/'test.xlsx',data_only=True)
        self.assertEqual(original['Breakdown ']['B14'].value,2)
        original.close()

    def test_edit_does_not_add_excel_notes_and_preview_keeps_provenance(self):
        report = ui.get_report(self.job)
        report['families'][0]['review_notes'] = {'Breakdown ': {'B14': 'Nguồn: input.xlsx\nSUM!W20'}}
        ui.write_json(self.folder/'review.json', report)
        self.edit('B14', '9')
        wb = openpyxl.load_workbook(self.folder/'files'/'test.xlsx')
        self.assertFalse(any(c.comment for s in wb for row in s for c in row))
        wb.close()
        preview = ui.preview(self.job, 0, 0)
        self.assertIn('SUM!W20', preview['rows'][13][1]['comment'])

    def test_fix_error_and_invalid_week_updates_release_and_cumulative(self):
        result=self.edit('C15','8')
        self.assertFalse(any(p['kind']=='error' for p in result['problems']))
        self.assertTrue(any(p['kind']=='week' for p in result['problems']))
        result=self.edit('A15','2028/01',revision=1)
        self.assertEqual(result['problems'],[])
        wb=self.values()
        self.assertEqual(wb['TEST']['D20'].value,20)
        self.assertEqual(wb['TEST']['A20'].value,'2028/01')
        self.assertEqual(wb['TEST']['E20'].value.date(),date(2028,1,3))
        self.assertEqual(wb['Release Qty']['A4'].value.date(),date(2028,1,3))
        result=self.edit('A15',None,revision=2,restore=True)
        self.assertTrue(any(p['kind']=='week' for p in result['problems']))

    def test_reject_formula_edit_nan_invalid_week_duplicate_and_stale_revision(self):
        for cell,value in [('D14','10'),('B14','NaN'),('B14','=1+1'),('A15','2027/53'),('A15','2026/53'),('B5','new identity')]:
            with self.subTest(cell=cell,value=value),self.assertRaises(ValueError):
                self.edit(cell,value)
        self.edit('B14','5')
        with self.assertRaisesRegex(ValueError,'cập nhật'):
            self.edit('B14','6',revision=0)

    def test_confirm_and_metadata_sync(self):
        self.edit('B9','250')
        wb=self.values()
        self.assertEqual(wb['TEST']['B14'].value,350)
        self.assertEqual(wb['TEST']['B16'].value,390)
        self.edit('B7','NEW',revision=1)
        self.assertEqual(self.values()['TEST']['B11'].value,'NEW / INTERNATIONAL')

    def test_preview_original_and_formula_locks(self):
        self.edit('B14','9')
        result=ui.preview(self.job,0,0)
        cell=result['rows'][13][1]
        self.assertEqual(cell['original'],2)
        self.assertEqual(cell['value'],9)
        self.assertTrue(cell['changed'])
        self.assertEqual(cell['editable'],'number')
        self.assertIsNone(result['rows'][13][3]['editable'])
        self.assertTrue(result['rows'][13][3]['formula'].startswith('=SUM'))

    def test_edit_preserves_template_formatting(self):
        path=self.folder/'files'/'test.xlsx'
        wb=openpyxl.load_workbook(path)
        cell=wb['Breakdown ']['B9']
        cell.fill=PatternFill('solid',fgColor='114433')
        cell.font=Font(name='Arial',size=14,bold=True,color='AA2200')
        cell.number_format='0.0000'
        ui.recalculate(wb)
        process.save_with_cache(wb,path)
        wb.close()
        self.edit('B9','250')
        wb=openpyxl.load_workbook(path)
        cell=wb['Breakdown ']['B9']
        self.assertEqual(cell.fill.fgColor.rgb,'00114433')
        self.assertEqual(cell.font.color.rgb,'00AA2200')
        self.assertEqual(cell.font.size,14)
        self.assertTrue(cell.font.bold)
        self.assertEqual(cell.number_format,'0.0000')
        wb.close()

    def test_text_is_literal_and_date_validation(self):
        self.edit('G19','=not a formula',sheet=1)
        wb=openpyxl.load_workbook(self.folder/'files'/'test.xlsx')
        self.assertEqual(wb['TEST']['G19'].data_type,'s')
        wb.close()
        with self.assertRaises(ValueError):
            self.edit('H19','not a date',sheet=1,revision=1)
        self.edit('H19','2026-10-05',sheet=1,revision=1)
        self.assertEqual(self.values()['TEST']['H19'].value.date(),date(2026,10,5))

    def test_batch_saves_once_to_new_session_and_keeps_original(self):
        old = (self.folder/'files/test.xlsx').read_bytes()
        with patch.object(ui, 'recalculate', wraps=ui.recalculate) as calc:
            result = ui.apply_batch(self.job, 0, dict(revision=0, edits=[
                dict(sheet=0, cell='B14', value=12), dict(sheet=0, cell='C14', value=8),
                dict(sheet=1, cell='G19', value='Updated')]))
        self.assertEqual(calc.call_count, 1)
        self.assertEqual((self.folder/'files/test.xlsx').read_bytes(), old)
        book = openpyxl.load_workbook(ui.directory(result['new_job_id'])/'files/test.xlsx', data_only=True)
        self.assertEqual(book['TEST']['B19'].value, 20)
        self.assertEqual(book['TEST']['G19'].value, 'Updated')
        book.close()

    def test_invalid_batch_does_not_publish_partial_edits(self):
        old = (self.folder/'files/test.xlsx').read_bytes()
        with self.assertRaises(ValueError):
            ui.apply_batch(self.job, 0, dict(revision=0, edits=[
                dict(sheet=0, cell='B14', value=12), dict(sheet=0, cell='C14', value='NaN')]))
        self.assertEqual((self.folder/'files/test.xlsx').read_bytes(), old)
        self.assertEqual(ui.get_report(self.job)['families'][0]['revision'], 0)
        other = next(p for p in ui.STORE.iterdir() if p.name != self.job)
        self.assertEqual(ui.read_job(other.name)['status'], 'error')

    def test_batch_week_swap_validates_final_state(self):
        self.edit('A15', '2027/01')
        result = ui.apply_batch(self.job, 0, dict(revision=1, edits=[
            dict(sheet=0, cell='A14', value='2027/01'),
            dict(sheet=0, cell='A15', value='2026/53')]))
        book = openpyxl.load_workbook(ui.directory(result['new_job_id'])/'files/test.xlsx', data_only=True)
        self.assertEqual(book['Breakdown ']['A14'].value, '2027/01')
        self.assertEqual(book['Breakdown ']['A15'].value, '2026/53')
        book.close()

    @patch.object(ui.excel_engine, 'capability', return_value={'available': False, 'message': 'Chưa kết nối Excel'})
    def test_source_is_read_only_without_worker(self, capability):
        source = openpyxl.Workbook()
        source.active.title = 'SUM'
        source['SUM']['A1'] = 10
        source.save(self.folder/'input.xlsx')
        shutil.copy2(self.folder/'input.xlsx', self.folder/'baseline'/'input.xlsx')
        report = ui.get_report(self.job)
        report['source_file'] = {
            'id': -1, 'file': 'input.xlsx', 'name': 'File nguồn: input.xlsx',
            'original_name': 'input.xlsx', 'revision': 0, 'edits': [],
            'is_source': True, 'changed_cells': 0, 'problems': []
        }
        ui.write_json(self.folder/'review.json', report)

        preview = ui.preview(self.job, -1, 0)
        self.assertIsNone(preview['rows'][0][0]['editable'])
        with self.assertRaisesRegex(ValueError, 'Chưa kết nối Excel'):
            ui.apply_edit(self.job, -1, {'sheet': 0, 'cell': 'A1', 'revision': 0, 'value': 20})
        unchanged = openpyxl.load_workbook(self.folder/'input.xlsx', data_only=True)
        self.assertEqual(unchanged['SUM']['A1'].value, 10)
        unchanged.close()

    def source_fixture(self):
        # Retain regression coverage for the old SUM override format, which full
        # workbook edits migrate before asking Excel to recalculate.
        route = patch.object(ui.source_edit, 'apply', side_effect=lambda server, job, report, entry, payload:
                             ui.apply_source_edit(job, report, entry, payload))
        route.start()
        self.addCleanup(route.stop)
        capability = patch.object(ui.excel_engine, 'capability', return_value={'available': True})
        capability.start()
        self.addCleanup(capability.stop)
        source = openpyxl.Workbook()
        source.active.title = 'SUM'
        source['SUM']['W16'] = '=SUM(1,1)'
        source._formula_cache = {('SUM', 'W16'): 2}
        process.save_with_cache(source, self.folder/'input.xlsx')
        source.close()
        shutil.copy2(self.folder/'input.xlsx', self.folder/'baseline'/'input.xlsx')
        report = ui.get_report(self.job)
        report['families'][0]['review_notes'] = {'Breakdown ': {'B14': 'Nguồn: input.xlsx\nSUM!W16\nQuantity'}}
        report['source_file'] = {'id': -1, 'file': 'input.xlsx', 'name': 'Source',
                                 'is_source': True, 'revision': 0, 'edits': [], 'problems': []}
        ui.write_json(self.folder/'review.json', report)

    def test_source_override_updates_consumers_preserves_cache_and_unrelated_edits(self):
        self.source_fixture()
        original = (self.folder/'input.xlsx').read_bytes()
        self.edit('C14', '8')
        preview = ui.preview(self.job, -1, 0)
        self.assertEqual(preview['rows'][15][22]['editable'], 'source_cell')
        result = ui.apply_edit(self.job, -1, dict(sheet=0, cell='W16', revision=0, value='12.5'))
        self.assertEqual(result['updated_families'], [0])
        wb = self.values()
        self.assertEqual(wb['Breakdown ']['B14'].value, 12.5)
        self.assertEqual(wb['Breakdown ']['C14'].value, 8)
        self.assertEqual(wb['TEST']['B19'].value, 20.5)
        self.assertEqual(wb['TEST']['D19'].value, 20.5)
        self.assertEqual((self.folder/'input.xlsx').read_bytes(), original)
        preview = ui.preview(self.job, -1, 0)
        self.assertEqual(preview['rows'][15][22]['value'], 12.5)
        self.assertEqual(preview['rows'][15][22]['original'], 2)
        with self.assertRaises(ValueError):
            ui.apply_edit(self.job, -1, dict(sheet=0, cell='W16', revision=0, value='9'))
        result = ui.apply_edit(self.job, -1, dict(sheet=0, cell='W16', revision=1, restore=True))
        self.assertEqual(result['changed_cells'], 0)
        self.assertEqual(self.values()['TEST']['B19'].value, 10)

    def test_source_edit_failure_keeps_outputs_and_report_unchanged(self):
        self.source_fixture()
        before = (self.folder/'files'/'test.xlsx').read_bytes()
        report = (self.folder/'review.json').read_bytes()
        with patch.object(ui, 'write_json', side_effect=OSError('disk full')):
            with self.assertRaises(OSError):
                ui.apply_edit(self.job, -1, dict(sheet=0, cell='W16', revision=0, value='9'))
        self.assertEqual((self.folder/'files'/'test.xlsx').read_bytes(), before)
        self.assertEqual((self.folder/'review.json').read_bytes(), report)
        for value in ('NaN', 'Infinity', '=SUM(1,1)'):
            with self.assertRaises(ValueError):
                ui.apply_edit(self.job, -1, dict(sheet=0, cell='W16', revision=0, value=value))

    def test_source_edit_updates_every_consumer_and_invalidates_old_family_revision(self):
        self.source_fixture()
        report = ui.get_report(self.job)
        second = json.loads(json.dumps(report['families'][0]))
        second.update(id=1, file='second.xlsx')
        report['families'].append(second)
        for subfolder in ('files', 'baseline'):
            shutil.copy2(self.folder/subfolder/'test.xlsx', self.folder/subfolder/'second.xlsx')
        ui.write_json(self.folder/'review.json', report)
        result = ui.apply_edit(self.job, -1, dict(sheet=0, cell='W16', revision=0, value=21))
        self.assertEqual(result['updated_families'], [0, 1])
        for family in ui.get_report(self.job)['families']:
            self.assertEqual(family['revision'], 1)
            book = openpyxl.load_workbook(self.folder/'files'/family['file'], data_only=True)
            self.assertEqual(book['TEST']['B19'].value, 24)
            book.close()
            with self.assertRaises(ValueError):
                ui.apply_edit(self.job, family['id'], dict(sheet=0, cell='B14', revision=0, value=1))


if __name__=='__main__':
    unittest.main()
