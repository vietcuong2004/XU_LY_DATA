import io
import os
from pathlib import Path
import tempfile
import threading
import unittest
from http.server import ThreadingHTTPServer
from unittest.mock import patch
from urllib.request import Request, urlopen
from urllib.error import HTTPError
from zipfile import ZipFile

import openpyxl
import excel_engine
import excel_worker


class WorkerTransportTests(unittest.TestCase):
    def setUp(self):
        self.server = ThreadingHTTPServer(('127.0.0.1', 0), excel_worker.Handler)
        self.thread = threading.Thread(target=self.server.serve_forever, daemon=True)
        self.thread.start()
        self.url = f'http://127.0.0.1:{self.server.server_port}'
        env = patch.dict(os.environ, {'SHIPMENT_CALCULATOR': 'excel', 'EXCEL_WORKER_URL': self.url, 'EXCEL_WORKER_TOKEN': 'private-test-token-at-least-32-chars'})
        env.start(); self.addCleanup(env.stop)
        self.addCleanup(self.stop)

    def stop(self):
        self.server.shutdown(); self.server.server_close(); self.thread.join(timeout=5)

    def test_unauthorized_request_never_starts_excel(self):
        with patch.object(excel_engine, 'edit_and_calculate_local') as engine:
            with self.assertRaises(HTTPError) as error:
                urlopen(Request(self.url+'/recalculate', data=b'invalid', method='POST'))
            self.assertEqual(error.exception.code, 401)
            engine.assert_not_called()

    def test_authenticated_roundtrip_and_worker_failure_preserves_source(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp)/'input.xlsx'
            book = openpyxl.Workbook(); book.active['A1']=1; book.save(path); book.close()
            before = path.read_bytes()
            seen = []
            def calculate(received, edits):
                seen.append((received.read_bytes(), edits))
                book = openpyxl.load_workbook(received)
                book.active['A1'] = 7; book.save(received); book.close()
            edits = [dict(sheet='Sheet', cell='A1', kind='number', value=7)]
            with patch.object(excel_engine, 'edit_and_calculate_local', side_effect=calculate):
                excel_engine.edit_and_calculate(path, edits)
            self.assertEqual(seen, [(before, edits)])
            book = openpyxl.load_workbook(path); self.assertEqual(book.active['A1'].value, 7); book.close()
            updated = path.read_bytes()
            with patch.object(excel_engine, 'edit_and_calculate_local', side_effect=ValueError('failure')):
                with self.assertRaises(ValueError):
                    excel_engine.edit_and_calculate(path, edits)
            self.assertEqual(path.read_bytes(), updated)


if __name__ == '__main__':
    unittest.main()
