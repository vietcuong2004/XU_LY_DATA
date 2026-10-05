"""Regression checks against the real template, including grouped column widths."""
from copy import copy
from pathlib import Path
from types import SimpleNamespace
import tempfile
import unittest

import openpyxl

import process
from template_format import apply_template_format,effective_column


class TemplateFormatTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        # Build a fresh fixture: user output folders may be removed at any time.
        cls.generated = tempfile.TemporaryDirectory()
        cls.addClassCleanup(cls.generated.cleanup)
        report = process.run(SimpleNamespace(
            input=str(process.ROOT/'templates/input_sample.xlsx'),
            template=str(process.find_default_template()), output=cls.generated.name,
            family=['BARBIE 2728'], season='2728', week=40, start_week=None,
            end_week=None, date=None, strict=False, overwrite=False))
        cls.output_path = Path(cls.generated.name)/report['families'][0]['file']

    def setUp(self):
        self.template = openpyxl.load_workbook(process.find_default_template())
        self.wb = openpyxl.load_workbook(self.output_path)
        self.addCleanup(self.template.close)
        self.addCleanup(self.wb.close)

    def test_exact_styles_dimensions_merges_and_views_at_template_size(self):
        self.wb.worksheets[0].delete_rows(84,1000)
        self.wb.worksheets[1].delete_rows(89,1000)
        apply_template_format(self.wb,self.template)
        checked = 0
        for sheet,reference in zip(self.wb.worksheets[:2],self.template.worksheets[:2]):
            self.assertEqual(set(map(str,sheet.merged_cells.ranges)),set(map(str,reference.merged_cells.ranges)))
            self.assertEqual(sheet.freeze_panes,reference.freeze_panes)
            self.assertEqual(sheet.sheet_view.showGridLines,reference.sheet_view.showGridLines)
            self.assertEqual(sheet.page_setup,reference.page_setup)
            for col in range(1,sheet.max_column+1):
                self.assertEqual(effective_column(sheet,col).width,effective_column(reference,col).width)
            for row in sheet:
                for cell in row:
                    original = reference[cell.coordinate]
                    for attr in ('font','fill','border','alignment','protection','number_format'):
                        self.assertEqual(copy(getattr(cell,attr)),copy(getattr(original,attr)),
                                         f'{sheet.title}!{cell.coordinate} {attr}')
                    checked += 1
            for row in range(1,sheet.max_row+1):
                a,b = sheet.row_dimensions.get(row),reference.row_dimensions.get(row)
                self.assertEqual(a.height if a else None,b.height if b else None)
        self.assertGreater(checked,2500)

    def test_preserves_values_and_grouped_widths_with_extended_weeks(self):
        before = {(s.title,c.coordinate):c.value for s in self.wb for row in s for c in row if c.value is not None}
        apply_template_format(self.wb,self.template)
        after = {(s.title,c.coordinate):c.value for s in self.wb for row in s for c in row if c.value is not None}
        self.assertEqual(before,after)
        b=self.wb['Breakdown ']
        self.assertEqual(effective_column(b,3).width,17.85546875)
        self.assertEqual(effective_column(b,15).width,15.7109375)
        self.assertEqual(copy(b['B129'].font),copy(self.template['Breakdown ']['B82'].font))
        self.assertEqual(copy(b['B130'].border),copy(self.template['Breakdown ']['B83'].border))
        self.assertEqual(copy(self.wb['Release Qty']['A4'].fill),copy(self.template['Release Qty']['A4'].fill))


if __name__=='__main__':
    unittest.main()
