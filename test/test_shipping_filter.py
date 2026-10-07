from datetime import date, datetime
from copy import copy
from pathlib import Path
from types import SimpleNamespace
import tempfile
import unittest

import openpyxl
import process


class ShippingFilterTests(unittest.TestCase):
    def test_positive_item_not_net_total_and_reselection_after_edit(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws['A16'] = '2026/40'
        ws['W16'], ws['X16'], ws['Y16'] = 5, -5, 0
        infos = [{'name': 'A', 'columns': [23, 24]}, {'name': 'B', 'columns': [25]}]
        reader = process.SheetReader(ws)
        weeks = process.extract_weeks(reader, [])
        chosen, report = process.filter_shipping_families(infos, reader, weeks, (2026, 40))
        self.assertEqual([f['name'] for f in chosen], ['A'])
        self.assertEqual(report['excluded'], ['B'])
        ws['W16'], ws['Y16'] = 0, 10
        chosen, _ = process.filter_shipping_families(infos, reader, weeks, (2026, 40))
        self.assertEqual([f['name'] for f in chosen], ['B'])
        ws['Y16'] = None
        self.assertEqual(process.filter_shipping_families(infos, reader, weeks, (2026, 40))[0], [])
        ws['Y16'] = '#REF!'
        with self.assertRaisesRegex(ValueError, 'B:'):
            process.filter_shipping_families(infos, reader, weeks, (2026, 40))
        with self.assertRaisesRegex(ValueError, 'Không có tuần'):
            process.filter_shipping_families(infos, reader, weeks, (2026, 41))

    def test_sample_exports_ten_families_with_full_schedule(self):
        source = process.ROOT/'2026 INTERNAL SCHEDULE FERRERO-WK40(LOG)-T.xlsx'
        with tempfile.TemporaryDirectory() as tmp:
            args = SimpleNamespace(input=source, template=None, output=tmp, family=None,
                                   season='2728', week=40, start_week=None, end_week=None,
                                   shipping_week=(2026, 40), date=date(2026, 10, 6),
                                   strict=False, overwrite=False)
            report = process.run(args)
            self.assertEqual(report['shipping_filter']['detected'], 26)
            self.assertEqual(len(report['families']), 10)
            self.assertEqual(report['weeks'], 117)
            self.assertEqual(len(list(Path(tmp).glob('*.xlsx'))), 10)
            from ui_server import recalculate, reconcile_source
            wb = openpyxl.load_workbook(source, data_only=True)
            template = openpyxl.load_workbook(process.find_default_template())
            reader = process.SheetReader(wb['SUM'])
            for entry in report['families']:
                output = Path(tmp)/entry['file']
                checks = reconcile_source(reader, output, entry)
                self.assertGreater(checks, 0)
                formulas = openpyxl.load_workbook(output, data_only=False)
                values = openpyxl.load_workbook(output, data_only=True)
                total_row = 14 + report['weeks']
                self.assertEqual(formulas['Breakdown '].cell(total_row, 1).value, 'Total')
                self.assertIs(formulas['Breakdown '].cell(total_row + 1, 2).value, False)
                self.assertIs(formulas['Breakdown '].cell(total_row + 2, 2).value, False)
                self.assertEqual(formulas['Breakdown '].cell(total_row + 3, 2).data_type, 'f')
                self.assertEqual(formulas['Breakdown '].cell(total_row + 5, 2).data_type, 'f')
                cached_check = values['Breakdown '].cell(total_row + 5, 2).value
                self.assertTrue(isinstance(cached_check, bool) or cached_check in process.ERRORS)
                self.assertEqual(copy(formulas['Breakdown '].cell(total_row, 2).font),
                                 copy(template['Breakdown ']['B84'].font))
                self.assertEqual(copy(formulas['Breakdown '].cell(total_row + 5, 2).fill),
                                 copy(template['Breakdown ']['B89'].fill))
                summary_total_row = 19 + report['weeks']
                summary = formulas[entry['sheet']]
                summary_values = values[entry['sheet']]
                self.assertEqual(summary.cell(summary_total_row, 1).value, 'TOTAL QTY: ')
                self.assertEqual(summary.cell(summary_total_row + 1, 1).value, 'EACH: ')
                self.assertEqual(summary.cell(summary_total_row, 2).data_type, 'f')
                self.assertEqual(summary.cell(summary_total_row + 1, 2).data_type, 'f')
                self.assertIsInstance(summary_values.cell(summary_total_row, entry['markets'] + 4).value,
                                      datetime)
                recalculate(formulas)
                recalculated_check = formulas._formula_cache[('Breakdown ', f'B{total_row + 5}')]
                self.assertTrue(isinstance(recalculated_check, bool) or recalculated_check in process.ERRORS)
                formulas.close()
                values.close()
            wb.close()
            template.close()


if __name__ == '__main__':
    unittest.main()
