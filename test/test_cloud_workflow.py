"""Real input -> all families -> Python source edit -> cold cloud restore."""
from pathlib import Path
import os
import tempfile
import unittest
from unittest.mock import patch

import openpyxl
import ui_server as ui
import cloud_store as cloud
from test.test_cloud_store import MemoryS3


class WorkflowTests(unittest.TestCase):
    def test_full_generation_edit_and_cold_restore(self):
        root=Path(__file__).resolve().parents[1]
        source=root/'templates/input_sample.xlsx'
        memory=MemoryS3()
        with tempfile.TemporaryDirectory() as temp, patch.object(ui,'STORE',Path(temp)/'worker1'), \
             patch.object(cloud,'client',return_value=memory), \
             patch.dict(os.environ,{'SHIPMENT_S3_BUCKET':'test','SHIPMENT_CALCULATOR':'python'}):
            ui.STORE.mkdir()
            events=[]
            job=ui.create_job({'input':{'name':'input.xlsx','raw':source.read_bytes()}},on_progress=events.append)
            self.assertEqual(job['status'],'ready',job.get('message'))
            report=ui.get_report(job['id'])
            self.assertEqual(len(report['families']),26)
            wb=openpyxl.load_workbook(ui.directory(job['id'])/'input.xlsx',read_only=True)
            sheet=wb.sheetnames.index('Bel');wb.close()
            with cloud.edit_lock(job['id']):
                result=ui.apply_batch(job['id'],-1,{'revision':0,'edits':[
                    {'sheet':sheet,'cell':'C16','value':'321.25','value_type':'number'}]})
            new_id=result['new_job_id']
            self.assertNotEqual(job['id'],new_id)
            # A fresh instance has no local source or outputs at all.
            with patch.object(ui,'STORE',Path(temp)/'worker2'):
                folder=ui.directory(new_id)
                updated=ui.get_report(new_id)
                self.assertEqual(len(updated['families']),26)
                wb=openpyxl.load_workbook(folder/'input.xlsx',data_only=True)
                self.assertEqual(wb['SUM']['W16'].value,321.25);wb.close()
                entry=next(f for f in updated['families'] if f['name']=='KS HL PLAYMOBIL DISNEY 2627')
                wb=openpyxl.load_workbook(folder/'files'/entry['file'],data_only=True)
                self.assertEqual(wb['Breakdown ']['B14'].value,321.25);wb.close()
                self.assertEqual(ui.read_job(job['id'])['superseded_by'],new_id)
                self.assertEqual((ui.directory(job['id'])/'input.xlsx').read_bytes(),source.read_bytes())


if __name__=='__main__':unittest.main()
