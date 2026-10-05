"""Full-source editing: publication, failures, and optional real Excel checks."""
from pathlib import Path
import os
import shutil
import tempfile
import unittest
from unittest.mock import patch

import openpyxl
import excel_engine
import process
import source_edit
import ui_server as ui


class FullSourceTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.store = Path(self.temp.name)
        patcher = patch.object(ui, 'STORE', self.store)
        patcher.start(); self.addCleanup(patcher.stop)
        patcher = patch.object(excel_engine, 'capability', return_value={'available': True})
        patcher.start(); self.addCleanup(patcher.stop)
        self.job = 'b' * 32
        self.folder = self.store/self.job
        (self.folder/'baseline').mkdir(parents=True)
        (self.folder/'files').mkdir()
        book = openpyxl.Workbook()
        book.active.title = 'Input'
        book['Input']['A1'] = 3
        book.create_sheet('SUM')['A1'] = '=Input!A1*2'
        book._formula_cache = {('SUM', 'A1'): 6}
        process.save_with_cache(book, self.folder/'input.xlsx')
        book.close()
        shutil.copy2(self.folder/'input.xlsx', self.folder/'baseline/input.xlsx')
        shutil.copy2(self.folder/'input.xlsx', self.folder/'template.xlsx')
        ui.write_json(self.folder/'job.json', dict(id=self.job, status='ready', season='2728', week=40, source='source.xlsx'))
        ui.write_json(self.folder/'review.json', dict(families=[], date='2026-10-06',
            source_file=dict(id=-1, name='source', file='input.xlsx', is_source=True, edit_mode='full', revision=0, edits=[])))

    def calculate(self, path, edits):
        book = openpyxl.load_workbook(path)
        for edit in edits:
            book[edit['sheet']][edit['cell']] = edit['value']
        book._formula_cache = {('SUM', 'A1'): book['Input']['A1'].value * 2}
        process.save_with_cache(book, path)
        book.close()

    def generate(self, job, args):
        folder = ui.directory(job)
        (folder/'baseline').mkdir()
        shutil.copy2(folder/'input.xlsx', folder/'baseline/input.xlsx')
        ui.write_json(folder/'review.json', dict(families=[{'id': 0}], source_file={'is_source': True, 'id': -1}))
        ui.update_job(job, status='ready')

    def test_new_generation_and_current_source_preview_original_restore(self):
        before = (self.folder/'input.xlsx').read_bytes()
        with patch.object(excel_engine, 'edit_and_calculate', side_effect=self.calculate), patch.object(ui, 'generate', side_effect=self.generate):
            result = ui.apply_edit(self.job, -1, dict(sheet=0, cell='A1', revision=0, value='8', value_type='number'))
        new_job = result['new_job_id']
        self.assertNotEqual(new_job, self.job)
        self.assertEqual((self.folder/'input.xlsx').read_bytes(), before)
        self.assertEqual(ui.read_job(self.job)['superseded_by'], new_job)
        book = openpyxl.load_workbook(ui.directory(new_job)/'input.xlsx', data_only=True)
        self.assertEqual(book['SUM']['A1'].value, 16)
        book.close()
        self.assertEqual((ui.directory(new_job)/'baseline/input.xlsx').read_bytes(), before)
        preview = ui.preview(new_job, -1, 0)
        self.assertEqual(preview['rows'][0][0]['value'], 8)
        self.assertEqual(preview['rows'][0][0]['original'], 3)
        with self.assertRaisesRegex(ValueError, 'phiên mới'):
            ui.apply_edit(self.job, -1, dict(sheet=0, cell='A1', revision=0, value='9', value_type='number'))
        with patch.object(excel_engine, 'edit_and_calculate', side_effect=self.calculate), patch.object(ui, 'generate', side_effect=self.generate):
            restored = ui.apply_edit(new_job, -1, dict(sheet=0, cell='A1', revision=1, restore=True))
        self.assertEqual(restored['changed_cells'], 0)
        book = openpyxl.load_workbook(ui.directory(restored['new_job_id'])/'input.xlsx', data_only=True)
        self.assertEqual(book['Input']['A1'].value, 3)
        self.assertEqual(book['SUM']['A1'].value, 6)
        book.close()

    def test_worker_failure_keeps_parent_untouched(self):
        source = (self.folder/'input.xlsx').read_bytes()
        report = (self.folder/'review.json').read_bytes()
        with patch.object(excel_engine, 'edit_and_calculate', side_effect=ValueError('worker offline')):
            with self.assertRaisesRegex(ValueError, 'worker offline'):
                ui.apply_edit(self.job, -1, dict(sheet=0, cell='A1', revision=0, value=8, value_type='number'))
        self.assertEqual((self.folder/'input.xlsx').read_bytes(), source)
        self.assertEqual((self.folder/'review.json').read_bytes(), report)
        self.assertNotIn('superseded_by', ui.read_job(self.job))

    def test_multiple_source_cells_use_one_excel_call_and_one_generation(self):
        with patch.object(excel_engine, 'edit_and_calculate', side_effect=self.calculate) as engine, \
             patch.object(ui, 'generate', side_effect=self.generate) as generate:
            result = ui.apply_batch(self.job, -1, dict(revision=0, edits=[
                dict(sheet=0, cell='A1', value=8, value_type='number'),
                dict(sheet=1, cell='A1', value='=Input!A1*2', value_type='formula')]))
        self.assertEqual(engine.call_count, 1)
        self.assertEqual(generate.call_count, 1)
        self.assertEqual(len(engine.call_args.args[1]), 2)
        self.assertEqual(len(result['edits']), 2)

    def test_bad_source_batch_is_validated_before_excel(self):
        with patch.object(excel_engine, 'edit_and_calculate') as engine:
            with self.assertRaises(ValueError):
                ui.apply_batch(self.job, -1, dict(revision=0, edits=[
                    dict(sheet=0, cell='A1', value=8, value_type='number'),
                    dict(sheet=1, cell='A1', value='NaN', value_type='number')]))
            engine.assert_not_called()

    def test_types_are_explicit_and_formula_restore_keeps_formula(self):
        self.assertEqual(source_edit.cell_edit('SUM', 'A1', '=1+1', 'text')['kind'], 'text')
        for raw in ('NaN', 'Infinity', '=1+1'):
            with self.assertRaises(ValueError):
                source_edit.cell_edit('SUM', 'A1', raw, 'number')
        for raw in ('abc', '1+1'):
            with self.assertRaises(ValueError):
                source_edit.cell_edit('SUM', 'A1', raw, 'formula')
        book = openpyxl.load_workbook(self.folder/'input.xlsx')
        edit = source_edit.original_edit('SUM', book['SUM']['A1'])
        self.assertEqual(edit['kind'], 'formula')
        self.assertEqual(edit['value'], '=Input!A1*2')
        book.close()


@unittest.skipUnless(os.environ.get('RUN_EXCEL_TESTS') == '1', 'Opt in to desktop Excel integration')
class DesktopExcelTests(unittest.TestCase):
    def test_cross_sheet_calculation_and_style(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp)/'input.xlsx'
            book = openpyxl.Workbook()
            book.active.title = 'Injection'
            book['Injection']['A1'] = 4
            book['Injection']['A1'].number_format = '0.000'
            book.create_sheet('Production')['A1'] = '=Injection!A1*2'
            book.create_sheet('SUM')['A1'] = '=Production!A1+1'
            book.save(path);book.close()
            excel_engine.edit_and_calculate_local(path, [dict(sheet='Injection', cell='A1', kind='number', value=9)])
            book = openpyxl.load_workbook(path, data_only=True)
            self.assertEqual(book['Production']['A1'].value, 18)
            self.assertEqual(book['SUM']['A1'].value, 19)
            self.assertEqual(book['Injection']['A1'].number_format, '0.000')
            book.close()

    def test_real_workbook_remote_worker_and_all_family_regeneration(self):
        import threading
        from http.server import ThreadingHTTPServer
        import excel_worker
        root = Path(__file__).resolve().parent.parent
        source = root/'2026 INTERNAL SCHEDULE FERRERO-WK40(LOG)-T.xlsx'
        if not source.exists():
            self.skipTest('Real source fixture is not available')
        server = ThreadingHTTPServer(('127.0.0.1', 0), excel_worker.Handler)
        thread = threading.Thread(target=server.serve_forever, daemon=True)
        thread.start()
        try:
            with tempfile.TemporaryDirectory(dir=root/'.qa') as temp, \
                 patch.object(ui, 'STORE', Path(temp)), \
                 patch.dict(os.environ, {'EXCEL_WORKER_URL': f'http://127.0.0.1:{server.server_port}',
                                         'EXCEL_WORKER_TOKEN': 'integration-test-secret-1234567890'}):
                job = 'd' * 32
                folder = Path(temp)/job
                (folder/'baseline').mkdir(parents=True)
                shutil.copy2(source, folder/'input.xlsx')
                shutil.copy2(source, folder/'baseline/input.xlsx')
                shutil.copy2(process.find_default_template(root), folder/'template.xlsx')
                ui.write_json(folder/'job.json', dict(id=job, status='ready', source=source.name, season='2728', week=40))
                ui.write_json(folder/'review.json', dict(families=[], date='2026-10-06',
                    source_file=dict(id=-1, name=source.name, file='input.xlsx', is_source=True, edit_mode='full', revision=0, edits=[])))
                book = openpyxl.load_workbook(source, read_only=True)
                bel = book.sheetnames.index('Bel')
                sheets = book.sheetnames
                book.close()
                result = ui.apply_edit(job, -1, dict(sheet=bel, cell='C16', revision=0, value=321.25, value_type='number'))
                updated = ui.directory(result['new_job_id'])
                book = openpyxl.load_workbook(updated/'input.xlsx', data_only=True)
                self.assertEqual(book.sheetnames, sheets)
                self.assertEqual(book['SUM']['W16'].value, 321.25)
                book.close()
                report = ui.get_report(result['new_job_id'])
                self.assertEqual(len(report['families']), 26)
                family = next(f for f in report['families'] if f['name']=='KS HL PLAYMOBIL DISNEY 2627')
                book = openpyxl.load_workbook(updated/'files'/family['file'], data_only=True)
                self.assertEqual(book['Breakdown ']['B14'].value, 321.25)
                book.close()
                self.assertEqual((folder/'input.xlsx').read_bytes(), source.read_bytes())
        finally:
            server.shutdown()
            server.server_close()
            thread.join(timeout=5)


if __name__ == '__main__':
    unittest.main()
