"""Unit tests và đối chiếu toàn bộ kết quả với SUM (python -m unittest -v)."""
from collections import OrderedDict
from datetime import date
from pathlib import Path
import json
import tempfile
import unittest
from unittest.mock import patch

import openpyxl

import process


class ExtractionTests(unittest.TestCase):
    def test_review_notes_are_separate_and_migration_preserves_user_notes(self):
        from openpyxl.comments import Comment
        wb = openpyxl.Workbook()
        sheet = wb.active
        process.review_note(sheet['A1'], 'Nguồn: input.xlsx\nSUM!W20')
        self.assertIsNone(sheet['A1'].comment)
        self.assertIn('SUM!W20', wb._review_notes[sheet.title]['A1'])
        sheet['B1'].comment = Comment('Nguồn: input.xlsx\nSUM!W20\nCache.', 'Source')
        sheet['C1'].comment = Comment('Tuần 2026/40.', 'Review')
        sheet['D1'].comment = Comment('Ghi chú riêng của tôi', 'Review')
        migrated = process.remove_generated_notes(wb)
        self.assertEqual(set(migrated[sheet.title]), {'B1', 'C1'})
        self.assertIsNone(sheet['B1'].comment)
        self.assertIsNone(sheet['C1'].comment)
        self.assertEqual(sheet['D1'].comment.text, 'Ghi chú riêng của tôi')

    def test_publish_uses_destination_sibling_and_preserves_bytes(self):
        with tempfile.TemporaryDirectory() as root:
            root = Path(root)
            private, public = root/'private', root/'public'
            private.mkdir()
            public.mkdir()
            source, destination = private/'file.xlsx', public/'file.xlsx'
            source.write_bytes(b'PK\x03\x04exact workbook bytes')
            destination.write_bytes(b'old version')
            original_replace = Path.replace
            promotions = []
            def inspect_replace(stage, target):
                promotions.append((stage.parent, Path(target).parent))
                return original_replace(stage, target)
            with patch.object(Path, 'replace', inspect_replace):
                process.publish_file(source, destination)
            self.assertEqual(destination.read_bytes(), source.read_bytes())
            self.assertEqual(promotions, [(public, public)])
            self.assertEqual(list(public.glob('*.tmp')), [])
            # Repairing an existing file must also preserve its exact bytes.
            process.publish_file(destination, destination)
            self.assertEqual(destination.read_bytes(), source.read_bytes())

    def test_merged_headers_noncontiguous_family_and_auxiliary(self):
        wb = openpyxl.Workbook()
        s = wb.active
        s['V3'] = 'CAPSULE'
        s['W3'] = 'A'
        s.merge_cells('W3:Y3')
        s['W5'], s['X5'], s['Y5'] = 'Item 1', 'Total', 'Nr. Pallet'
        s['Z3'], s['Z5'] = 'B', 'Item 2'
        s['AA3'], s['AA5'] = 'A', 'Item 3'
        s['AB5'] = 'No family: must not inherit'
        families = process.extract_families(s)
        self.assertEqual([f['name'] for f in families], ['A', 'B'])
        self.assertEqual(families[0]['ranges'], [[23, 25], [27, 27]])
        self.assertEqual(families[0]['columns'], [23, 27])

    def test_missing_formula_cache_is_not_zero(self):
        values, formulas = openpyxl.Workbook(), openpyxl.Workbook()
        formulas.active['W3'] = '=Other!A1'
        reader = process.SheetReader(values.active, formulas.active)
        with self.assertRaisesRegex(ValueError, 'cache'):
            reader.value(3, 23)

    def test_iso_boundary_invalid_week_and_duplicate(self):
        wb = openpyxl.Workbook()
        s = wb.active
        s['A16'], s['A17'], s['A18'] = '2026/53', '2027/01', '2027/53'
        issues = []
        weeks = process.extract_weeks(process.SheetReader(s), issues)
        self.assertEqual(weeks[0]['monday'], date(2026, 12, 28))
        self.assertEqual(weeks[1]['monday'], date(2027, 1, 4))
        self.assertIsNone(weeks[2]['monday'])
        self.assertEqual(len(issues), 1)
        s['A19'] = '2027/01'
        with self.assertRaisesRegex(ValueError, 'lặp'):
            process.extract_weeks(process.SheetReader(s), [])

    def test_output_cache_error_propagation_and_literal_text(self):
        wb = openpyxl.Workbook()
        s = wb.active
        wb._formula_cache = {}
        process.put(s, 1, 1, 12.5)
        process.put(s, 2, 1, -2)
        process.sum_formula(s, 3, 1, [s['A1'], s['A2']])
        process.put(s, 4, 1, '#REF!')
        process.sum_formula(s, 5, 1, [s['A3'], s['A4']])
        process.put(s, 6, 1, '=literal PO')
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / 'test.xlsx'
            process.save_with_cache(wb, path)
            values = openpyxl.load_workbook(path, data_only=True)
            formulas = openpyxl.load_workbook(path)
            self.assertEqual(values.active['A3'].value, 10.5)
            self.assertEqual(values.active['A5'].value, '#REF!')
            self.assertEqual(formulas.active['A3'].data_type, 'f')
            self.assertEqual(formulas.active['A6'].data_type, 's')
            values.close()
            formulas.close()

    def test_safe_names(self):
        self.assertEqual(process.safe_filename('A/B:C?'), 'A_B_C_')
        self.assertLessEqual(len(process.sheet_name('a' * 70)), 31)
        self.assertNotEqual(process.sheet_name('Release Qty'), 'Release Qty')
        self.assertEqual(process.safe_filename('CON'), '_CON')


class RealWorkbookTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.root = Path(__file__).parent
        report_path = cls.root / 'output_families' / 'validation_report.json'
        if not report_path.exists():
            raise unittest.SkipTest('Chạy process.py với file nguồn trước để kiểm tra tích hợp.')
        cls.report = json.loads(report_path.read_text(encoding='utf-8'))
        cls.source = openpyxl.load_workbook(cls.report['input'], data_only=True)
        cls.sheet = cls.source['SUM']
        cls.merged = {}
        for area in cls.sheet.merged_cells.ranges:
            for row in range(area.min_row, min(15, area.max_row) + 1):
                for col in range(area.min_col, area.max_col + 1):
                    cls.merged[row, col] = cls.sheet.cell(area.min_row, area.min_col).value
        cls.week_rows = {cls.sheet.cell(r, 1).value: r for r in range(16, cls.sheet.max_row + 1)
                         if isinstance(cls.sheet.cell(r, 1).value, str) and '/' in cls.sheet.cell(r, 1).value}

    @classmethod
    def tearDownClass(cls):
        cls.source.close()

    def value(self, row, col):
        return self.merged.get((row, col), self.sheet.cell(row, col).value)

    def assert_value(self, actual, expected):
        if isinstance(expected, (int, float)):
            self.assertAlmostEqual(actual, expected, places=7)
        else:
            self.assertEqual(actual, expected)

    def total(self, values):
        error = next((v for v in values if isinstance(v, str) and v.startswith('#')), None)
        return error or sum(v for v in values if isinstance(v, (int, float)))

    def test_all_output_cells_against_source(self):
        for entry in self.report['families']:
            with self.subTest(family=entry['name']):
                path = self.root / 'output_families' / entry['file']
                values = openpyxl.load_workbook(path, data_only=True)
                formulas = openpyxl.load_workbook(path)
                self.assertEqual(values.sheetnames, ['Breakdown ', entry['sheet'], 'Release Qty'])
                b, s, release = values.worksheets
                self.assertEqual(s['E2'].value, None)
                self.assertEqual(s['E4'].value, entry['name'])
                markets = OrderedDict()
                for col in entry['source_columns']:
                    key = tuple(str(self.value(r, col) or '').strip() for r in (4, 15, 2))
                    markets.setdefault(key, []).append(col)
                source_cols = [c for group in markets.values() for c in group]
                for out_col, source_col in enumerate(source_cols, 2):
                    for out_row, source_row in [(9,11), (10,12), (12,8), (13,9)]:
                        self.assert_value(b.cell(out_row, out_col).value, self.value(source_row, source_col))
                unique_items = list(dict.fromkeys((str(self.value(6,c) or '').strip(),
                                                   str(self.value(5,c) or '').strip()) for c in source_cols))
                running = 0
                for out_row in range(14, b.max_row + 1):
                    label = b.cell(out_row, 1).value
                    src_row = self.week_rows[label]
                    row_values = [self.sheet.cell(src_row, c).value for c in source_cols]
                    for out_col, expected in enumerate(row_values, 2):
                        self.assert_value(b.cell(out_row, out_col).value, expected)
                    for out_col, key in enumerate(unique_items, len(source_cols) + 2):
                        expected = self.total([self.sheet.cell(src_row, c).value for c in source_cols
                                               if (str(self.value(6,c) or '').strip(), str(self.value(5,c) or '').strip()) == key])
                        self.assert_value(b.cell(out_row, out_col).value, expected)
                    summary_row = out_row + 5
                    for out_col, group in enumerate(markets.values(), 2):
                        expected = self.total([self.sheet.cell(src_row, c).value for c in group])
                        self.assert_value(s.cell(summary_row, out_col).value, expected)
                    total_col = len(markets) + 2
                    expected_total = self.total(row_values)
                    self.assert_value(s.cell(summary_row, total_col).value, expected_total)
                    running = self.total([running, expected_total])
                    self.assert_value(s.cell(summary_row, total_col + 1).value, running)
                    year, week = map(int, label.split('/'))
                    try:
                        monday = date.fromisocalendar(year, week, 1)
                    except ValueError:
                        monday = None
                    actual = s.cell(summary_row, total_col + 2).value
                    self.assertEqual(actual.date() if actual else None, monday)
                    actual = release.cell(out_row - 11, 1).value
                    self.assertEqual(actual.date() if actual else None, monday)
                self.assertFalse(formulas._external_links)
                for sh in formulas:
                    for row in sh:
                        for cell in row:
                            if cell.data_type == 'f':
                                self.assertNotIn('#REF!', cell.value)
                                self.assertNotIn('[', cell.value)
                                self.assertIsNotNone(values[sh.title][cell.coordinate].value)
                values.close()
                formulas.close()


if __name__ == '__main__':
    unittest.main()
