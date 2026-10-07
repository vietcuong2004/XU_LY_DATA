"""Tách SUM thành workbook theo Family; bố cục theo mẫu BARBIE.

Không tính lại công thức nguồn: dùng kết quả Excel đã lưu (data_only).
Các công thức đầu ra được tạo mới, kèm cache để xem ngay khi chưa mở Excel.
"""
from __future__ import annotations

import argparse
from collections import OrderedDict
from copy import copy
from datetime import date, datetime, timedelta, timezone
import json
import math
from pathlib import Path
import re
import shutil
import sys
import tempfile
import uuid
from zipfile import ZipFile, ZIP_DEFLATED
import xml.etree.ElementTree as ET

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as letter
from openpyxl.utils.datetime import to_excel
from openpyxl.workbook.properties import CalcProperties
import pandas as pd
from template_format import apply_template_format
ROOT = Path(__file__).resolve().parent


def find_default_template(parent: Path | None = None) -> Path:
    filename = 'BARBIE 2728 _Weekly shipment schedule 2728_WK39.xlsx'
    base = parent or ROOT
    candidates = [
        base / 'templates' / filename,
        base / 'tempalates' / filename,
        base / filename,
        ROOT / 'templates' / filename,
        ROOT / 'tempalates' / filename,
        ROOT / filename,
    ]
    for c in candidates:
        if c.is_file():
            return c
    return ROOT / 'templates' / filename


DEFAULT_TEMPLATE = 'templates/BARBIE 2728 _Weekly shipment schedule 2728_WK39.xlsx'
ERRORS = {'#REF!', '#VALUE!', '#N/A', '#DIV/0!', '#NAME?', '#NUM!', '#NULL!'}
NUMBER_FORMAT = '#,##0.###;[Red]-#,##0.###;0'
WEEK_RE = re.compile(r'^(\d{4})/(\d{1,2})$')
LOGISTICS = ['Total            (weekly)', 'CUM (Total)', 'Date', 'NR. OF WEEK',
             'ASN NO.', 'ETD', 'ETA', 'MEANS OF TRANSPORT', 'SHIPPING LINE',
             'VESSEL NAME', 'NOTE', 'CONTAINER NO.', 'SEAL NO.', 'INVOICE NO.']


def clean(value):
    return '' if value is None else str(value).strip()


class SheetReader:
    """Chỉ lan truyền giá trị trong đúng phạm vi ô gộp, không fill-forward."""
    def __init__(self, sheet, formula_sheet=None, issues=None):
        self.sheet = sheet
        self.formula_sheet = formula_sheet
        self.issues = issues if issues is not None else []
        self.anchors = {}
        self.reported = set()
        for area in sheet.merged_cells.ranges:
            for row in range(area.min_row, min(area.max_row, 15) + 1):
                for col in range(area.min_col, area.max_col + 1):
                    self.anchors[row, col] = (area.min_row, area.min_col)

    def value(self, row, col):
        row, col = self.anchors.get((row, col), (row, col))
        cell = self.sheet.cell(row, col)
        value = cell.value
        if self.formula_sheet is not None:
            source = self.formula_sheet.cell(row, col)
            if source.data_type == 'f' and value is None:
                raise ValueError(f'{self.sheet.title}!{cell.coordinate}: công thức chưa có cache. '
                                 'Mở nguồn bằng Excel, Calculate Now, lưu lại rồi chạy tool.')
        if cell.data_type == 'f':
            raise ValueError('extract_families cần sheet được mở với data_only=True')
        if cell.data_type == 'e' and cell.coordinate not in self.reported:
            self.issues.append({'type': 'source_error', 'cell': cell.coordinate, 'value': value})
            self.reported.add(cell.coordinate)
        return value


def extract_families(sum_sheet):
    """Trả về các Family, các dải liên tiếp và cột item; gộp cả dải rời nhau."""
    reader = sum_sheet if isinstance(sum_sheet, SheetReader) else SheetReader(sum_sheet)
    families = OrderedDict()
    previous = None
    for col in range(22, reader.sheet.max_column + 1):
        name = clean(reader.value(3, col))
        if not name or name.upper() in {'CAPSULE', 'BLUE-OPTION'}:
            previous = None
            continue
        if name in ERRORS:
            raise ValueError(f'Tên Family bị lỗi ở cột {letter(col)}: {name}')
        family = families.setdefault(name, {'name': name, 'ranges': [], 'columns': []})
        if previous == name:
            family['ranges'][-1][1] = col
        else:
            family['ranges'].append([col, col])
        previous = name
        item = clean(reader.value(5, col))
        if item and item.casefold() not in {'total', 'nr. pallet', 'nr.pallet'}:
            family['columns'].append(col)
    return [dict(f, start_column=f['ranges'][0][0], end_column=f['ranges'][-1][1])
            for f in families.values() if f['columns']]


def extract_weeks(reader, issues, start=None, end=None):
    weeks, seen = [], set()
    for row in range(16, reader.sheet.max_row + 1):
        labels = [clean(reader.value(row, col)) for col in (1, 22)]
        matches = [WEEK_RE.fullmatch(v) for v in labels]
        keys = {(int(m[1]), int(m[2])) for m in matches if m}
        if not keys:
            continue
        if len(keys) != 1:
            raise ValueError(f'Nhẫn tuần A{row} và V{row} không khớp: {labels}')
        year, week = key = keys.pop()
        if start and key < start or end and key > end:
            continue
        label = f'{year}/{week:02d}'
        if label in seen:
            raise ValueError(f'Tuần bị lặp: {label} tại dòng {row}')
        seen.add(label)
        try:
            monday = date.fromisocalendar(year, week, 1)
        except ValueError:
            monday = None
            issues.append({'type': 'invalid_iso_week', 'cell': f'A{row}', 'value': label,
                           'action': 'Giữ nhãn và số lượng; để trống ngày, không tự đổi tuần.'})
        weeks.append({'label': label, 'row': row, 'week': week, 'monday': monday})
    if not weeks:
        raise ValueError('Không tìm thấy tuần giao hàng trong phạm vi đã chọn.')
    return weeks


def filter_shipping_families(infos, reader, weeks, shipping_week):
    label = f'{shipping_week[0]}/{shipping_week[1]:02d}'
    date.fromisocalendar(*shipping_week, 1)
    target = next((w for w in weeks if w['label'] == label), None)
    if target is None:
        raise ValueError(f'Không có tuần {label} trong SUM. Hãy chọn tuần có trong file nguồn.')
    selected, excluded = [], []
    for info in infos:
        quantities = [numeric(reader.value(target['row'], c), f'{letter(c)}{target["row"]}')
                      for c in info['columns']]
        if any(isinstance(q, (int, float)) and q > 0 for q in quantities):
            selected.append(info)
        elif any(q in ERRORS for q in quantities if isinstance(q, str)):
            raise ValueError(f'{info["name"]}: số liệu tuần {label} có lỗi Excel; chưa thể xác định có giao hàng.')
        else:
            excluded.append(info['name'])
    return selected, {'week': label, 'detected': len(infos), 'selected': len(selected),
                      'excluded': excluded, 'rule': 'Có ít nhất một item có số lượng > 0; giữ toàn bộ lịch và item.'}


def numeric(value, coordinate):
    if value is None or value == '':
        return None
    if isinstance(value, str) and value in ERRORS:
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value):
        return value
    raise ValueError(f'{coordinate}: số lượng không hợp lệ: {value!r}')


def prepare_family(info, reader, weeks, template, report_date, source_name):
    items = []
    for col in info['columns']:
        get = lambda r: reader.value(r, col)
        destination = clean(get(15))
        match = re.match(r'[A-Za-z0-9]{4}', destination)
        item = {'column': col, 'name': clean(get(5)), 'mpg': clean(get(6)),
                'country': clean(get(4)), 'version': clean(get(2)), 'po': get(7),
                'code': get(8), 'capsule': get(9), 'destination': destination,
                'dest_code': match[0] if match else destination,
                'confirm': numeric(get(11), f'{letter(col)}11'),
                'option': numeric(get(12), f'{letter(col)}12'),
                'quantities': [numeric(get(w['row']), f'{letter(col)}{w["row"]}') for w in weeks]}
        if not item['country']:
            raise ValueError(f'{info["name"]}: thiếu thị trường tại {letter(col)}4')
        # Giữ riêng các phiên bản (ví dụ Italy INTERNATIONAL và ISRAEL).
        item['market'] = (item['country'], destination, item['version'])
        items.append(item)
    markets = OrderedDict()
    for item in items:
        markets.setdefault(item['market'], []).append(item)
    ordered_items = [item for group in markets.values() for item in group]
    for col, item in enumerate(ordered_items, 2):
        item['output_column'] = col
    return dict(info, items=ordered_items, markets=markets, weeks=weeks, template=template,
                report_date=report_date, source=source_name,
                kind='KJ' if re.match(r'^KJ\b', info['name'], re.I) else 'KS')


def style_cell(target, source):
    for attr in ('font', 'fill', 'border', 'alignment', 'protection'):
        setattr(target, attr, copy(getattr(source, attr)))
    target.number_format = source.number_format


def setup_sheet(wb, title, reference, rows, columns, row_map, col_map):
    sheet = wb.create_sheet(title)
    for row in range(1, rows + 1):
        ref_row = row_map(row)
        sheet.row_dimensions[row].height = reference.row_dimensions[ref_row].height or 24
        for col in range(1, columns + 1):
            style_cell(sheet.cell(row, col), reference.cell(ref_row, col_map(col)))
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = reference.sheet_view.zoomScale or 75
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_options.horizontalCentered = True
    sheet.print_area = f'A1:{letter(columns)}{rows}'
    return sheet


def put(sheet, row, col, value):
    cell = sheet.cell(row, col)
    cell.value = value
    if isinstance(value, str) and value not in ERRORS:
        cell.data_type = 's'  # Mã PO/nhãn bắt đầu '=' vẫn là dữ liệu, không là công thức.
    return cell


def cached_value(cell):
    if cell.data_type == 'f':
        return cell.parent.parent._formula_cache[(cell.parent.title, cell.coordinate)]
    return cell.value


def sum_value(cells):
    values = [cached_value(c) for c in cells]
    error = next((v for v in values if isinstance(v, str) and v in ERRORS), None)
    return error or sum(v for v in values if isinstance(v, (int, float)))


def formula(sheet, row, col, expression, value):
    cell = sheet.cell(row, col, expression)
    sheet.parent._formula_cache[(sheet.title, cell.coordinate)] = value
    cell.number_format = NUMBER_FORMAT
    return cell


def sum_formula(sheet, row, col, cells):
    cells = list(cells)
    refs = [f"'{c.parent.title.replace(chr(39), chr(39)*2)}'!{c.coordinate}"
            if c.parent is not sheet else c.coordinate for c in cells]
    return formula(sheet, row, col, '=SUM(' + ','.join(refs) + ')', sum_value(cells))


def range_sum_formula(sheet, row, col, start_row, end_row, source_col=None):
    source_col = source_col or col
    cells = [sheet.cell(source_row, source_col) for source_row in range(start_row, end_row + 1)]
    column = letter(source_col)
    return formula(sheet, row, col, f'=SUM({column}{start_row}:{column}{end_row})', sum_value(cells))


def review_note(cell, text):
    """Keep provenance in the review report, without Excel's red Note markers."""
    wb = cell.parent.parent
    if not hasattr(wb, '_review_notes'):
        wb._review_notes = {}
    wb._review_notes.setdefault(cell.parent.title, {})[cell.coordinate] = text


def remove_generated_notes(wb):
    """Migrate only notes previously authored by this tool; preserve other notes."""
    notes = {}
    for sheet in wb:
        for row in sheet:
            for cell in row:
                note = cell.comment
                if not note:
                    continue
                generated = (
                    note.author == 'Source' and note.text.startswith('Nguồn: ') and '\nSUM!' in note.text
                    or note.author == 'Validation' and (
                        note.text.startswith('Tuần ') or 'không hợp lệ theo ISO; cần kiểm tra nguồn.' in note.text)
                    or note.author == 'Review' and re.fullmatch(r'Tuần \d{4}/\d{2}\.', note.text)
                )
                if generated:
                    notes.setdefault(sheet.title, {})[cell.coordinate] = note.text
                    cell.comment = None
    return notes


def mark_source(cell, source_cell, data):
    review_note(cell, f'Nguồn: {data["source"]}\nSUM!{source_cell.coordinate}\n'
                'Giá trị lấy từ cache đã lưu của file nguồn.')
    # Giữ BOLD/RED của kế hoạch nguồn, không mang công thức liên kết sang workbook mới.
    font = copy(cell.font)
    font.bold = source_cell.font.bold
    font.color = copy(source_cell.font.color)
    cell.font = font
    if cell.data_type == 'e':
        cell.fill = PatternFill('solid', fgColor='FFC7CE')


def merge_label(sheet, row, first, last):
    if last > first:
        sheet.merge_cells(start_row=row, start_column=first, end_row=row, end_column=last)


def finish_widths(sheet, label_width=30):
    for col in range(1, sheet.max_column + 1):
        # Tính theo cỡ font mẫu (nhiều ô là Arial 14/18, lớn hơn font mặc định Excel).
        widths = [len(str(c.value)) * (c.font.sz or 11) / 11 * 1.05 + 3
                  for column in sheet.iter_cols(min_col=col, max_col=col) for c in column
                  if c.value is not None and c.data_type != 'f']
        sheet.column_dimensions[letter(col)].width = label_width if col == 1 else min(40, max([20, *widths]))
    for row in sheet:
        for cell in row:
            alignment = copy(cell.alignment)
            alignment.wrap_text = True
            alignment.vertical = 'center'
            cell.alignment = alignment


def build_breakdown_sheet(wb, family_data):
    d = family_data
    groups = OrderedDict()
    for item in d['items']:
        groups.setdefault((item['mpg'], item['name']), []).append(item)
    last_col = 1 + len(d['items']) + len(groups)
    last_week_row = 13 + len(d['weeks'])
    footer_rows = 6
    sheet = setup_sheet(wb, 'Breakdown ', d['template']['Breakdown '], last_week_row + footer_rows,
                        last_col, lambda r: min(r, 14),
                        lambda c: 1 if c == 1 else (14 if c > 1 + len(d['items']) else 2))
    labels = {4: 'Dest. Code', 5: 'Item', 6: 'MPG', 7: 'Version: ', 8: 'PCS/carton',
              9: 'confirm', 10: 'Option', 11: 'Total', 12: 'Industrial Code: ', 13: 'capsule color'}
    for row, label in labels.items():
        put(sheet, row, 1, label)
    put(sheet, 2, 2, d['kind'])
    merge_label(sheet, 2, 2, last_col)
    for group in d['markets'].values():
        first, last = group[0]['output_column'], group[-1]['output_column']
        put(sheet, 3, first, group[0]['country'])
        put(sheet, 4, first, group[0]['dest_code'])
        for row in (3, 4):
            merge_label(sheet, row, first, last)
    for item in d['items']:
        col = item['output_column']
        for row, key in {5:'name', 6:'mpg', 7:'version', 9:'confirm', 10:'option',
                         12:'code', 13:'capsule'}.items():
            cell = put(sheet, row, col, item[key])
            cell.number_format = NUMBER_FORMAT if row in (9, 10) else 'General'
        sum_formula(sheet, 11, col, [sheet.cell(9, col), sheet.cell(10, col)])
        for row, value, week in zip(range(14, last_week_row + 1), item['quantities'], d['weeks']):
            cell = put(sheet, row, col, value)
            cell.number_format = NUMBER_FORMAT
            mark_source(cell, d['source_sheet'].cell(week['row'], item['column']), d)
    total_start = 2 + len(d['items'])
    put(sheet, 3, total_start, 'TOTAL ' + d['kind'])
    merge_label(sheet, 3, total_start, last_col)
    merge_label(sheet, 4, total_start, last_col)
    for col, ((mpg, name), items) in enumerate(groups.items(), total_start):
        put(sheet, 5, col, name)
        put(sheet, 6, col, mpg)
        for row in [9, 10, 11, *range(14, last_week_row + 1)]:
            sum_formula(sheet, row, col, [sheet.cell(row, i['output_column']) for i in items])
    for row, week in enumerate(d['weeks'], 14):
        put(sheet, row, 1, week['label'])

    # Preserve the six-row reconciliation footer from the template below the
    # dynamically-sized week block. Formulas and cached values are generated
    # for this Family, so they remain useful before Excel recalculates the file.
    total_row = last_week_row + 1
    put(sheet, total_row, 1, 'Total')
    for col in range(2, last_col + 1):
        range_sum_formula(sheet, total_row, col, 14, last_week_row)
    reference = d['template']['Breakdown ']
    reference_week_end = next(
        row - 1 for row in range(14, reference.max_row + 2)
        if not WEEK_RE.fullmatch(clean(reference.cell(row, 1).value)))
    reference_total_start = next(
        col for col in range(2, reference.max_column + 1)
        if clean(reference.cell(3, col).value).upper().startswith('TOTAL '))
    for row in (total_row + 1, total_row + 2):
        for col in range(2, total_start):
            reference_col = min(col, reference_total_start - 1)
            reference_row = reference_week_end + 1 + (row - total_row)
            put(sheet, row, col, reference.cell(reference_row, reference_col).value)
    difference_row = total_row + 3
    tolerance_row = total_row + 4
    check_row = total_row + 5
    reference_tolerance_row = reference_week_end + 5
    for col in range(2, total_start):
        total_value = cached_value(sheet.cell(total_row, col))
        released_value = cached_value(sheet.cell(11, col))
        difference = (total_value - released_value
                      if isinstance(total_value, (int, float)) and isinstance(released_value, (int, float))
                      else total_value if isinstance(total_value, str) and total_value in ERRORS
                      else released_value if isinstance(released_value, str) and released_value in ERRORS
                      else None)
        formula(sheet, difference_row, col,
                f'={letter(col)}{total_row}-{letter(col)}11', difference)
        reference_col = min(col, reference_total_start - 1)
        tolerance = reference.cell(reference_tolerance_row, reference_col).value
        tolerance = tolerance if isinstance(tolerance, (int, float)) else 15.6
        put(sheet, tolerance_row, col, tolerance).number_format = NUMBER_FORMAT
        passed = (difference if isinstance(difference, str) and difference in ERRORS
                  else difference <= tolerance if isinstance(difference, (int, float))
                  else False)
        formula(sheet, check_row, col,
                f'=+{letter(col)}{difference_row}<={letter(col)}{tolerance_row}', passed)
    sheet.freeze_panes = 'B14'
    sheet.print_title_rows = '1:13'
    finish_widths(sheet)
    for row in (5, 7, 12, 13):
        sheet.row_dimensions[row].height = max(sheet.row_dimensions[row].height or 0, 44)
    return sheet


def sheet_name(name):
    name = re.sub(r'[\\/*?:\[\]]', '_', name).strip("'")[:31] or 'Family'
    if name.casefold() in {'breakdown ', 'release qty'}:
        name = ('Family ' + name)[:31]
    return name


def joined(items, key):
    return ' / '.join(dict.fromkeys(clean(i[key]) for i in items if clean(i[key])))


def build_summary_sheet(wb, family_data):
    d = family_data
    markets = list(d['markets'].values())
    n = len(markets)
    total_col = n + 2
    last_week_row = 18 + len(d['weeks'])
    footer_rows = 3
    sheet = setup_sheet(wb, sheet_name(d['name']), d['template'].worksheets[1],
                        last_week_row + footer_rows, n + 15,
                        lambda r: min(r, 19),
                        lambda c: c if c == 1 else (2 if c <= n + 1 else c - n + 3))
    # Header có màu/chữ riêng từng cột; không lặp style cột thị trường B lên D/E.
    reference = d['template'].worksheets[1]
    for row in range(1, 9):
        for col in range(1, sheet.max_column + 1):
            style_cell(sheet.cell(row, col), reference.cell(row, min(col, reference.max_column)))
    # Header giống mẫu: dòng 3-6; bảng thị trường 9-18; tuần từ dòng 19.
    for row, a, b, label, value in [
        (3, 'Past Shipments', 'BOLD', 'Supplier :', 'NIIGATA'),
        (4, 'Changes', 'RED', 'Name of Family :', d['name']),
        (5, 'Train', None, 'Requested Weekly Output :', None),
        (6, 'Air', None, 'Date:', d['report_date'])]:
        for col, text in [(1,a), (2,b), (4,label), (5,value)]:
            put(sheet, row, col, text)
        merge_label(sheet, row, 5, 8)
        sheet.row_dimensions[row].height = 42
    sheet['E6'].number_format = 'yyyy-mm-dd'
    sheet['B3'].font = Font(name='Arial', size=14, bold=True)
    sheet['B4'].font = Font(name='Arial', size=14, color='FF0000')
    labels = {10:'Dest. Code', 11:'Version: ', 12:'Capsule: ', 13:'No. of items',
              14:'Released Qty: ', 15:'Option Qty: ', 16:'TOTAL QTY: ', 17:'PO No.', 18:'Format'}
    for row, text in labels.items():
        put(sheet, row, 1, text)
    breakdown = wb['Breakdown ']
    for col, group in enumerate(markets, 2):
        for row, value in {9:group[0]['country'], 10:group[0]['dest_code'],
                           11:group[0]['version'], 12:joined(group, 'capsule'),
                           13:len({(i['mpg'],i['name']) for i in group}),
                           17:joined(group, 'po'), 18:'T00'}.items():
            put(sheet, row, col, value).number_format = 'General'
        for row, source_row in [(14,9), (15,10), (16,11)]:
            sum_formula(sheet, row, col, [breakdown.cell(source_row, i['output_column']) for i in group])
        for row in range(19, last_week_row + 1):
            sum_formula(sheet, row, col, [breakdown.cell(row - 5, i['output_column']) for i in group])
    for col, label in enumerate(LOGISTICS, total_col):
        put(sheet, 9, col, label)
        sheet.merge_cells(start_row=9, start_column=col, end_row=18, end_column=col)
    for row, week in enumerate(d['weeks'], 19):
        put(sheet, row, 1, week['label'])
        sum_formula(sheet, row, total_col, [sheet.cell(row, c) for c in range(2, total_col)])
        previous = [] if row == 19 else [sheet.cell(row - 1, total_col + 1)]
        sum_formula(sheet, row, total_col + 1, previous + [sheet.cell(row, total_col)])
        cell = put(sheet, row, total_col + 2, week['monday'])
        cell.number_format = 'yyyy-mm-dd'
        if week['monday'] is None:
            review_note(cell, f'{week["label"]} không hợp lệ theo ISO; cần kiểm tra nguồn.')
            cell.fill = PatternFill('solid', fgColor='FFF2CC')
        put(sheet, row, total_col + 3, week['week']).number_format = '0'

    # Preserve the three calculation rows below the schedule from the sample.
    total_row, each_row, continuation_row = last_week_row + 1, last_week_row + 2, last_week_row + 3
    put(sheet, total_row, 1, 'TOTAL QTY: ')
    put(sheet, each_row, 1, 'EACH: ')
    for col in range(2, n + 2):
        range_sum_formula(sheet, total_row, col, 19, last_week_row)
        numerator = cached_value(sheet.cell(total_row, col))
        divisor = cached_value(sheet.cell(13, col))
        error = next((v for v in (numerator, divisor) if isinstance(v, str) and v in ERRORS), None)
        if error:
            each_value = error
        elif not isinstance(divisor, (int, float)) or divisor == 0:
            each_value = '#DIV/0!'
        else:
            each_value = numerator / divisor
        formula(sheet, each_row, col,
                f'={letter(col)}{total_row}/{letter(col)}13', each_value)
    total_start, total_end = letter(2), letter(n + 1)
    formula(sheet, total_row, total_col, f'=SUM({total_start}{total_row}:{total_end}{total_row})',
            sum_value(sheet.cell(total_row, col) for col in range(2, n + 2)))
    formula(sheet, each_row, total_col, f'=SUM({total_start}{each_row}:{total_end}{each_row})',
            sum_value(sheet.cell(each_row, col) for col in range(2, n + 2)))
    range_sum_formula(sheet, continuation_row, total_col, 19, last_week_row)

    def add_formula(row, col, left, right):
        left_value, right_value = cached_value(left), cached_value(right)
        error = next((v for v in (left_value, right_value) if isinstance(v, str) and v in ERRORS), None)
        value = error or sum(v for v in (left_value, right_value) if isinstance(v, (int, float)))
        return formula(sheet, row, col, f'={left.coordinate}+{right.coordinate}', value)

    # The sample starts TOTAL QTY from the cumulative value three rows before
    # the last week, then continues EACH and the final helper row sequentially.
    add_formula(total_row, total_col + 1,
                sheet.cell(total_row, total_col), sheet.cell(max(19, last_week_row - 3), total_col + 1))
    add_formula(each_row, total_col + 1,
                sheet.cell(each_row, total_col), sheet.cell(total_row, total_col + 1))
    add_formula(continuation_row, total_col + 1,
                sheet.cell(continuation_row, total_col), sheet.cell(each_row, total_col + 1))

    date_col, week_col = total_col + 2, total_col + 3
    date_source_rows = (max(19, last_week_row - 3), total_row, each_row)
    for row, source_row in zip((total_row, each_row, continuation_row), date_source_rows):
        source_date = cached_value(sheet.cell(source_row, date_col))
        next_date = source_date + timedelta(days=7) if isinstance(source_date, (date, datetime)) else None
        cell = formula(sheet, row, date_col, f'={letter(date_col)}{source_row}+7', next_date)
        cell.number_format = 'yyyy-mm-dd'
        iso_week = next_date.isocalendar().week if next_date else None
        formula(sheet, row, week_col, f'=_xlfn.ISOWEEKNUM({letter(date_col)}{row})', iso_week).number_format = '0'
    sheet.freeze_panes = 'B19'
    sheet.print_title_rows = '1:18'
    finish_widths(sheet)
    sheet.column_dimensions['D'].width = max(sheet.column_dimensions['D'].width, 32)
    for row in (9, 11, 12, 17):
        sheet.row_dimensions[row].height = 44
    return sheet


def build_release_qty_sheet(wb, family_data=None):
    d = family_data or wb._family_data
    sheet = wb.create_sheet('Release Qty')
    # Danh sách ngày tuần hiện tại, không sao chép cut-off cũ trong mẫu.
    for row, week in enumerate(d['weeks'], 3):
        cell = put(sheet, row, 1, week['monday'])
        style_cell(cell, d['template']['Release Qty']['A3'])
        cell.number_format = 'yyyy-mm-dd'
        review_note(cell, f'Tuần {week["label"]}; ngày thứ Hai theo ISO.' if week['monday'] else
                    f'Tuần {week["label"]} không hợp lệ theo ISO; ngày để trống.')
        sheet.row_dimensions[row].height = 25
    sheet.column_dimensions['A'].width = 23
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = 'A3'
    sheet.print_area = f'A1:A{len(d["weeks"])+2}'
    return sheet


def publish_file(source, destination):
    """Publish atomically with the destination folder's inherited permissions.

    Moving a file out of TemporaryDirectory on Windows preserves its private
    ACL. A regular file created beside the destination inherits the intended
    folder permissions, so desktop Excel can open it under the user's account.
    Copy bytes only: copy2/copystat must not carry private file attributes.
    """
    destination = Path(destination)
    destination.parent.mkdir(parents=True, exist_ok=True)
    stage = destination.parent / f'.shipment-{uuid.uuid4().hex}.tmp'
    try:
        with stage.open('xb') as output, Path(source).open('rb') as input_file:
            shutil.copyfileobj(input_file, output)
        stage.replace(destination)
    finally:
        stage.unlink(missing_ok=True)


def save_with_cache(wb, path):
    """Ghi cache cho đúng các công thức SUM đã tính, vẫn giữ công thức chỉnh sửa được."""
    ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    ET.register_namespace('', ns['s'])
    path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(dir=path.parent) as temp:
        raw, final = Path(temp)/'raw.xlsx', Path(temp)/'final.xlsx'
        wb.save(raw)
        with ZipFile(raw) as source, ZipFile(final, 'w', ZIP_DEFLATED) as target:
            for entry in source.infolist():
                data = source.read(entry.filename)
                match = re.fullmatch(r'xl/worksheets/sheet(\d+)\.xml', entry.filename)
                if match:
                    title = wb.worksheets[int(match[1])-1].title
                    root = ET.fromstring(data)
                    for cell in root.findall('.//s:c', ns):
                        key = (title, cell.attrib['r'])
                        if key not in wb._formula_cache:
                            continue
                        value = wb._formula_cache[key]
                        cached = cell.find('s:v', ns)
                        if cached is None:
                            cached = ET.SubElement(cell, '{'+ns['s']+'}v')
                        if isinstance(value, bool):
                            cached.text = '1' if value else '0'
                            cell.set('t', 'b')
                        elif isinstance(value, (date, datetime)):
                            cached.text = str(to_excel(value))
                            cell.set('t', 'n')
                        else:
                            cached.text = str(value)
                            cell.set('t', 'e' if isinstance(value, str) and value in ERRORS else 'n')
                    data = ET.tostring(root, encoding='utf-8', xml_declaration=True)
                target.writestr(entry, data)
        publish_file(final, path)


def safe_filename(name):
    name = re.sub(r'[\\/:*?"<>|\x00-\x1f]', '_', name).strip(' .')[:130] or 'Family'
    if re.match(r'^(CON|PRN|AUX|NUL|COM[1-9]|LPT[1-9])(?:\.|$)', name, re.I):
        name = '_' + name
    return name


def week_arg(value):
    match = WEEK_RE.fullmatch(value)
    if not match or not 1 <= int(match[2]) <= 53:
        raise argparse.ArgumentTypeError('Dùng YYYY/WW (01..53).')
    return int(match[1]), int(match[2])


def run(args):
    source_path = Path(args.input).resolve()
    template_path = Path(args.template).resolve() if args.template else find_default_template(source_path.parent)
    output = Path(args.output).resolve()
    if not template_path.is_file():
        raise ValueError(f'Không thấy mẫu: {template_path}. Dùng --template để chỉ định.')
    if args.start_week and args.end_week and args.start_week > args.end_week:
        raise ValueError('--start-week phải <= --end-week')
    issues = []
    values = openpyxl.load_workbook(source_path, data_only=True)
    formulas = openpyxl.load_workbook(source_path, data_only=False)
    template = openpyxl.load_workbook(template_path)
    if 'SUM' not in values:
        raise ValueError('File nguồn không có sheet SUM.')
    if 'Breakdown ' not in template or 'Release Qty' not in template or len(template.worksheets) < 3:
        raise ValueError('File mẫu phải có Breakdown , bảng Family và Release Qty.')
    reader = SheetReader(values['SUM'], formulas['SUM'], issues)
    infos = extract_families(reader)
    if args.family:
        wanted = {n.casefold() for n in args.family}
        missing = wanted - {i['name'].casefold() for i in infos}
        if missing:
            raise ValueError('Không tìm thấy Family: ' + ', '.join(sorted(missing)))
        infos = [i for i in infos if i['name'].casefold() in wanted]
    if not infos:
        raise ValueError('Không tìm thấy Family có item.')
    weeks = extract_weeks(reader, issues, args.start_week, args.end_week)
    shipping_filter = None
    if getattr(args, 'shipping_week', None):
        infos, shipping_filter = filter_shipping_families(
            infos, reader, extract_weeks(reader, issues), args.shipping_week)
    stamp = args.date or datetime.now(timezone(timedelta(hours=7))).date()
    datasets = [prepare_family(i, reader, weeks, template, stamp, getattr(args, 'source_name', source_path.name)) for i in infos]
    output.mkdir(parents=True, exist_ok=True)
    report = {'input': str(source_path), 'template': str(template_path), 'date': stamp.isoformat(),
              'weeks': len(weeks), 'first_week': weeks[0]['label'], 'last_week': weeks[-1]['label'],
              'issues': issues, 'families': [], 'shipping_filter': shipping_filter,
              'notes': ['Dùng cache nguồn: cần tính lại và lưu Excel trước khi chạy nếu nguồn thay đổi.',
                        'Giữ nguyên đơn vị nguồn, không nhân 1000 và không làm tròn.',
                        'Tách thị trường theo quốc gia + destination + version.']}
    if args.strict and issues:
        (output/'validation_report.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
        raise ValueError('Strict: nguồn có vấn đề; xem validation_report.json. Chưa xuất workbook.')
    paths, used = [], set()
    for d in datasets:
        filename = safe_filename(d['name']) + f' _Weekly shipment schedule {args.season}_WK{args.week:02d}.xlsx'
        if filename.casefold() in used:
            raise ValueError(f'Tên file trùng sau làm sạch: {filename}')
        used.add(filename.casefold())
        path = output / filename
        if path in {source_path, template_path}:
            raise ValueError('Đường dẫn xuất trùng nguồn hoặc mẫu.')
        if path.exists() and not args.overwrite:
            raise ValueError(f'File đã có: {path.name}. Dùng --overwrite để ghi đè.')
        paths.append(path)
    for d, path in zip(datasets, paths):
        d['source_sheet'] = formulas['SUM']
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        wb._formula_cache = {}
        wb._family_data = d
        wb.calculation = CalcProperties(calcId=191029, fullCalcOnLoad=True, forceFullCalc=True)
        wb.loaded_theme = template.loaded_theme
        build_breakdown_sheet(wb, d)
        build_summary_sheet(wb, d)
        build_release_qty_sheet(wb)
        apply_template_format(wb, template)
        save_with_cache(wb, path)
        # Pandas kiểm tra độc lập tổng số lượng hợp lệ; lỗi được đếm riêng, không đổi thành 0.
        frame = pd.DataFrame([{'week': w['label'], 'quantity': q}
                              for i in d['items'] for w, q in zip(weeks, i['quantities'])])
        nums = pd.to_numeric(frame['quantity'], errors='coerce')
        error_count = sum(isinstance(v, str) and v in ERRORS for v in frame['quantity'])
        report['families'].append({'name': d['name'], 'file': path.name,
                                   'sheet': wb.worksheets[1].title, 'ranges': d['ranges'],
                                   'items': len(d['items']), 'markets': len(d['markets']),
                                   'source_columns': d['columns'], 'source_error_cells': error_count,
                                   'review_notes': getattr(wb, '_review_notes', {}),
                                   'valid_numeric_quantity_subtotal': float(nums.sum()),
                                   'complete_total': None if error_count else float(nums.sum())})
        print(f'OK: {d["name"]} ({len(d["items"])} items, {len(d["markets"])} markets)')
        if getattr(args, 'on_progress', None):
            args.on_progress(len(report['families']), len(datasets), d['name'])
    (output/'validation_report.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    print(f'Exported {len(paths)} files; {len(issues)} source issues. Report: {output / "validation_report.json"}')
    values.close()
    formulas.close()
    template.close()
    return report


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input', required=True, help='File Master .xlsx chứa SUM')
    parser.add_argument('--template', help='File mẫu BARBIE; mặc định tìm cạnh file nguồn')
    parser.add_argument('--output', default='output_families')
    parser.add_argument('--family', action='append', help='Chỉ xuất Family này; có thể lặp lại')
    parser.add_argument('--season', default='2728')
    parser.add_argument('--week', type=int, choices=range(1,54), default=40)
    parser.add_argument('--date', type=date.fromisoformat, help='Ngày báo cáo YYYY-MM-DD; mặc định UTC+7')
    parser.add_argument('--start-week', type=week_arg)
    parser.add_argument('--shipping-week', type=week_arg, help='Chỉ xuất Family có giao trong YYYY/WW; giữ toàn bộ lịch')
    parser.add_argument('--end-week', type=week_arg)
    parser.add_argument('--strict', action='store_true', help='Dừng nếu gặp lỗi nguồn/tuần ISO không hợp lệ')
    parser.add_argument('--overwrite', action='store_true')
    args = parser.parse_args()
    if not re.fullmatch(r'[A-Za-z0-9_-]+', args.season):
        parser.error('--season chỉ gồm chữ, số, _ hoặc -')
    try:
        run(args)
    except (ValueError, OSError, KeyError) as exc:
        print(f'ERROR: {exc}', file=sys.stderr)
        return 1
    return 0


if __name__ == '__main__':
    sys.exit(main())
